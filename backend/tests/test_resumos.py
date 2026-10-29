import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app import config, resumos as rs

BRUTO = Path(__file__).parent / "fixtures" / "relatorio_bruto_nov_2026.xlsx"
WHATS = """Encaminhada
Permissionário
12:58
Encaminhada
Janeiro a março 3.050,00
12:58
Abril 2025 a março 2026 3.350,00 12:58
Abril 2026 a julho 2026 3.850,00
Agosto e setembro 2026 4.084,00
"""


@pytest.fixture()
def rx(client, monkeypatch):
    monkeypatch.setattr(config, "RAIOX_SENHA", "secreta123")
    return {"X-Raiox-Token": client.post("/raiox/login", json={"senha": "secreta123"}).json()["token"]}


def test_valor_br():
    assert rs.valor_br("R$ 1.132,11") == 1132.11
    assert rs.valor_br("R$ 0,00") == 0 and rs.valor_br(None) == 0 and rs.valor_br("3350") == 3350
    assert rs.valor_br("3.350") == 3350 and rs.valor_br(12.5) == 12.5


def test_le_relatorio_bruto():
    r = rs.ler_relatorio(BRUTO.read_bytes(), "novembro.xlsx")
    assert r["id"] == "2026-11" and len(r["linhas"]) == 10
    assert all(isinstance(l["debito"], float) for l in r["linhas"])
    assert (r["credito"], r["debito"]) == (255.88, 4483.08)  # mesmos totais do rodapé do sistema


def test_texto_do_whatsapp():
    f = rs.ler_texto(WHATS)["faixas"]["permissionario"]
    assert [(x["de"], x["ate"], x["valor"]) for x in f] == [
        ("2025-01", "2025-03", 3050.0), ("2025-04", "2026-03", 3350.0), ("2026-04", "2026-07", 3850.0),
        ("2026-08", "2026-08", 4084.0), ("2026-09", "2026-09", 4084.0)]
    assert rs.custo_do_mes(f, "2025-11") == 3350 and rs.custo_do_mes(f, "2026-10") is None


def test_formatos_de_linha():
    t = rs.ler_texto("3350,00 - 01/25\n02/2025 3.400\na partir de outubro 2026 4.100,00", "motoristas")["faixas"]
    assert t["motorista1"] == t["motorista2"]
    assert [(x["de"], x["ate"], x["valor"]) for x in t["motorista1"]] == [
        ("2025-01", "2025-01", 3350.0), ("2025-02", "2025-02", 3400.0), ("2026-10", None, 4100.0)]
    assert rs.ler_texto("Motorista 1\nnov/25 3.514,73")["faixas"]["motorista1"][0]["de"] == "2025-11"
    assert rs.ler_texto("dezembro 2025 a fevereiro 3.000")["faixas"] == {"motorista1": [], "motorista2": [], "permissionario": []}


def test_planilha_de_custos():
    wb = Workbook()
    ws = wb.active
    ws.append(["Permissionário", "Motorista 1"])
    ws.append(["3350,00 - 01/25", "3514,73 - 01/25"])
    ws.append(["3350,00 - 02/25", "3514,73 - 02/25"])
    b = io.BytesIO(); wb.save(b)
    f = rs.ler_planilha_custos(b.getvalue())["faixas"]
    assert [x["de"] for x in f["permissionario"]] == ["2025-01", "2025-02"] and f["motorista1"][1]["valor"] == 3514.73
    assert f["motorista2"] == []


def test_api_fluxo_completo(client, rx):
    assert client.get("/resumos/custos").status_code == 403
    lido = client.post("/resumos/custos/texto", json={"texto": WHATS}, headers=rx).json()
    custos = {"permissionario": lido["faixas"]["permissionario"] + [{"de": "2026-10", "ate": None, "valor": 4100}],
              "motorista1": [{"de": "2025-01", "ate": None, "valor": 3514.73}], "motorista2": [{"de": "2025-01", "ate": None, "valor": 3514.73}]}
    assert client.put("/resumos/custos", json=custos, headers=rx).status_code == 200
    arqs = [("arquivos", ("novembro.xlsx", BRUTO.read_bytes()))]
    p = client.post("/resumos/previa", files=arqs, headers=rx).json()
    m = p["meses"][0]
    assert m["permissionario"] == 4100 and m["motorista1"] == 3514.73 and not m["avisos"]
    assert m["saldo"] == round(m["credito"] - m["debito"] - 4100 - 2 * 3514.73, 2)
    r = client.post("/resumos/gerar", files=arqs * 2, headers=rx)
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert sorted(z.namelist()) == ["novembro_26.xlsx", "novembro_26_2.xlsx", "resumo_nov26.xlsx"]
    ws = load_workbook(io.BytesIO(z.read("novembro_26.xlsx"))).active
    assert [c.value for c in ws[2]] == ["Cap", "Titulo", "Descrição", "Data de Pagamento", "Valor Crédito", "Valor Débito"]
    assert isinstance(ws["F3"].value, (int, float)) and ws["F3"].number_format != "General"


def test_gerar_pdf(client, rx):
    arqs = [("arquivos", ("novembro.xlsx", BRUTO.read_bytes()))]
    z = zipfile.ZipFile(io.BytesIO(client.post("/resumos/gerar", files=arqs, data={"formato": "ambos"}, headers=rx).content))
    assert sorted(z.namelist()) == ["novembro_26.pdf", "novembro_26.xlsx", "resumo_nov26.pdf", "resumo_nov26.xlsx"]
    assert z.read("novembro_26.pdf")[:4] == b"%PDF"
    assert client.post("/resumos/gerar", files=arqs, data={"formato": "doc"}, headers=rx).status_code == 422
