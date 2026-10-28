from pathlib import Path

import pytest

PDF = Path(__file__).parent / "fixtures" / "folha_oficina.pdf"


def enviar(client, centro="ADM04", rota="/folha/processar", **extra):
    dados = {"centro_custo": centro, **extra} if centro else extra
    return client.post(rota, files={"arquivo": ("folha.pdf", PDF.read_bytes(), "application/pdf")}, data=dados)


@pytest.fixture()
def resp(client):
    r = enviar(client)
    assert r.status_code == 200, r.text
    return r.json()


def por_desc(resp):
    return {l["descricao"]: l for l in resp["linhas"]}


def test_cabecalho_e_periodo(resp):
    assert resp["departamento_pdf"] == "OFICINA" and resp["centro_custo"]["codigo"] == "ADM04"
    assert resp["periodo"] == {"inicio": "2026-05-01", "fim": "2026-05-31"} and resp["avisos"] == []


def test_evento_do_departamento_usa_o_centro_escolhido(resp):
    sal = por_desc(resp)["Salário Base"]
    assert (sal["valor"], sal["codigo_centro_de_custo"], sal["codigo_item_despesa"]) == (54244.65, "ADM04", "*070")
    assert sal["titulo"] == "Folha Oficina" and sal["observacoes"] == "Folha maio 26 Oficina"
    assert sal["data_inicial"] == sal["data_liquidacao"] == "2026-05-30"


def test_evento_fixo_vai_para_ctpas_com_item_sem_prefixo(resp):
    inss = por_desc(resp)["INSS Folha"]
    assert (inss["valor"], inss["codigo_centro_de_custo"], inss["codigo_item_despesa"], inss["tipo_lancamento"]) \
        == (5542.99, "CTPAS", "501", "CREDITO")


def test_por_padrao_traz_as_49_linhas_como_o_excel(resp):
    assert len(resp["linhas"]) == 49
    d = por_desc(resp)
    assert d["Saldo de Salário"]["valor"] == 0 and d["Plano Odontológico"]["valor"] == 0


def test_ocultar_zerados(client):
    linhas = enviar(client, ocultar_zerados="true").json()["linhas"]
    assert 0 < len(linhas) < 49 and all(l["valor"] != 0 for l in linhas)


def test_conferencias_batem_no_pdf_da_oficina(resp):
    conf = {c["nome"]: c for c in resp["conferencias"]}
    assert all(c["ok"] for c in conf.values()), conf
    assert conf["Total de adicionais"]["pdf"] == 91656.18 and conf["GPS Total Líquido"]["pdf"] == 27870.18
    assert len(conf) == 6


def test_calculados(resp):
    d = por_desc(resp)
    assert d["PIS Folha"]["valor"] == d["PIS Folha - Oficina"]["valor"] == 702.18
    assert d["Total FGTS  MÊS"]["valor"] == d["Total FGTS apurado recibos s/CS"]["valor"] == 5892.58
    assert d["Plano de Saúde - Total "]["valor"] == 1350.84  # Plano De Saúde - Mens. (675.42) x 2
    assert d["TOTAL LÍQUIDO A PAGAR SEM FÉRIAS"]["valor"] == 32387.41  # líquido - adiantamento - líquido de férias
    assert d["Total Líquido INSS - Oficina"]["valor"] == 21677.82  # soma dos componentes da GPS
    assert d["Desconto Empréstimo Crédito Trabalhador"]["valor"] == 0  # Total - Ativos


def test_ordem_segue_o_csv_original(resp):
    ordem = [l["descricao"] for l in resp["linhas"]]
    assert ordem[0] == "Salário Base" and ordem.index("INSS Folha") < ordem.index("Empresa Funcionários")


def test_sem_centro_sugere_pelo_departamento_do_pdf(client):
    r = enviar(client, centro=None)
    assert r.status_code == 200 and r.json()["centro_custo"]["codigo"] == "ADM04"


def test_centro_invalido(client):
    assert enviar(client, centro="CTPAS").status_code == 200  # CTPAS também pode ser escolhido
    assert enviar(client, centro="XYZ99").status_code == 422
    assert enviar(client, centro="XXX").status_code == 422


def test_pdf_invalido(client):
    r = client.post("/folha/processar", files={"arquivo": ("x.pdf", b"nao e pdf", "application/pdf")},
                    data={"centro_custo": "ADM04"})
    assert r.status_code == 422


def test_evento_novo_gera_alerta_e_some_depois_de_cadastrar(client, db):
    from app.models import Evento
    from sqlalchemy import select
    ev = db.scalar(select(Evento).where(Evento.nome_pdf == "Hora Extra 50%"))
    ev.nome_pdf = "Hora Extra 50% (renomeada)"  # simula evento que o banco não conhece
    db.commit()
    novos = enviar(client).json()["eventos_novos"]
    assert [n["nome_pdf"] for n in novos] == ["Hora Extra 50%"] and novos[0]["codigo_pdf"] == "401"
    ev.nome_pdf = "Hora Extra 50%"
    db.commit()
    assert enviar(client).json()["eventos_novos"] == []


def test_csv_tem_as_17_colunas_na_ordem(client):
    r = enviar(client, rota="/folha/csv")
    assert r.status_code == 200 and "lancamentos_Oficina_05-2026.csv" in r.headers["content-disposition"]
    linhas = r.content.decode("utf-8-sig").splitlines()
    assert linhas[0] == ("titulo;descricao;tipo_lancamento;codigo_centro_de_custo;numero_conta;codigo_item_despesa;"
                         "codigo_fonte_recurso;id_forma_pagamento;num_parcelas;data_inicial;dias_intervalo;valor;"
                         "observacoes;id_lancamento_grupo;status_pagamento;data_liquidacao;banco_liquidacao")
    assert linhas[1] == ("Folha Oficina;Salário Base;DEBITO;ADM04; 1-1;*070;;;;30/05/2026;;54244,65;"
                         "Folha maio 26 Oficina;;PAGO;30/05/2026; 1-1")


ADM = Path(__file__).parent / "fixtures" / "folha_administrativo_08.pdf"
# valores do CSV de agosto que o Excel/usuário esperava (Administrativo, ADM01)
ESPERADO_ADM = {
    "Salário Base": 13610.31, "Férias": 3686.65, "Periculosidade 30%": 494.16, "Adicional Noturno 20%": 327.36,
    "Repouso Remunerado": 62.96, "1/3 Férias": 1393.61, "Salário Família": 67.54, "Vale Transporte": 223.85,
    "Plano De Saúde": 225.14, "Plano De Saúde - Mens.": 1350.84, "Plano Odontológico": 45.0,
    "Plano Odontológico - Mens.": 180.0, "INSS Férias": 453.05, "INSS Folha": 1090.05,
    "TOTAL LÍQUIDO A PAGAR (FÉRIAS)": 4851.23, "TOTAL LÍQUIDO A PAGAR SEM FÉRIAS": 5802.03,
    "Empresa Funcionários": 3915.01, "RAT Emp (RAT x FAP = 5,0721 %)": 992.86, "Terceiros 5,80 %": 1135.35,
    "PIS Folha": 195.75, "Total Líquido INSS - Oficina": 6043.22, "PIS Folha - Oficina": 195.75,
    "Total FGTS apurado recibos s/CS": 1468.71, "Total FGTS  MÊS": 1468.71, "Plano de Saúde - Total ": 2701.68,
}


def test_administrativo_agosto(client):
    r = client.post("/folha/processar", files={"arquivo": ("a.pdf", ADM.read_bytes(), "application/pdf")})
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert corpo["centro_custo"]["codigo"] == "ADM01" and corpo["avisos"] == []
    assert corpo["eventos_novos"] == []  # "Periculosidade" é apelido de "Periculosidade 30%"
    assert all(c["ok"] for c in corpo["conferencias"]), corpo["conferencias"]
    linhas = {l["descricao"]: l["valor"] for l in corpo["linhas"] if l["valor"]}
    # "Plano de Saúde - Total " aparece 2x (débito/crédito): o dict guarda o valor, que é o mesmo
    assert linhas == ESPERADO_ADM
    assert len(corpo["linhas"]) == 49


def test_irrf_ferias_do_bloco_darf_nao_vira_desconto(resp):
    # no PDF da Oficina o 'IRRF Férias' só existe no bloco DARF IR, não é desconto da folha
    assert "IRRF Férias" in por_desc(resp) and por_desc(resp)["IRRF Férias"]["valor"] == 0


def test_vincular_apelido(client):
    from sqlalchemy import select  # noqa: F401
    eventos = client.get("/eventos", params={"q": "Hora Extra 50%"}).json()
    r = client.post(f"/eventos/{eventos[0]['id']}/apelidos", json={"nome": "Hora Extra 60%"})
    assert r.status_code == 201 and r.json()["apelidos"] == ["Hora Extra 60%"]
    assert client.post(f"/eventos/{eventos[0]['id']}/apelidos", json={"nome": "hora extra 60%"}).status_code == 409


@pytest.mark.parametrize("arq,centro,liquido_sem_ferias", [
    ("folha_coleta_08.pdf", "ADM02", 327.20),      # sem 'TOTAL LÍQUIDO (FÉRIAS)' no PDF
    ("folha_oficina_08.pdf", "ADM04", 30717.48),
    ("folha_posto_08.pdf", "ADM03", 1936.96),      # 'Periculosidade 30%' com o nome completo
])
def test_outros_departamentos_de_agosto(client, arq, centro, liquido_sem_ferias):
    pdf = Path(__file__).parent / "fixtures" / arq
    r = client.post("/folha/processar", files={"arquivo": (arq, pdf.read_bytes(), "application/pdf")})
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert corpo["centro_custo"]["codigo"] == centro and corpo["avisos"] == [] and corpo["eventos_novos"] == []
    assert all(c["ok"] for c in corpo["conferencias"]), corpo["conferencias"]
    assert len(corpo["linhas"]) == 49
    d = {l["descricao"]: l for l in corpo["linhas"]}
    assert d["TOTAL LÍQUIDO A PAGAR SEM FÉRIAS"]["valor"] == liquido_sem_ferias
    assert d["Salário Base"]["codigo_centro_de_custo"] == centro


def test_socios_conselho(client):
    pdf = Path(__file__).parent / "fixtures" / "folha_socios_08.pdf"
    r = client.post("/folha/processar", files={"arquivo": ("s.pdf", pdf.read_bytes(), "application/pdf")})
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert corpo["perfil"] == "conselho" and corpo["centro_custo"]["codigo"] == "ADM01"
    assert corpo["eventos_novos"] == [] and corpo["avisos"] == []
    assert all(c["ok"] for c in corpo["conferencias"]), corpo["conferencias"]
    l = [(x["descricao"], x["tipo_lancamento"], x["codigo_centro_de_custo"], x["codigo_item_despesa"], x["valor"])
         for x in corpo["linhas"]]
    assert l == [
        ("INSS  Conselheiros/Diretoria", "DEBITO", "ADM01", "*056", 713.24),
        ("INSS  Conselheiros/Diretoria", "CREDITO", "CTPAS", "501", 713.24),
        ("INSS  Conselheiros/Diretoria - Empresa", "DEBITO", "ADM01", "*056", 1945.2),
        ("INSS  Conselheiros/Diretoria - Empresa", "CREDITO", "CTPAS", "501", 1945.2),
        ("IRRF Sobre Folha", "DEBITO", "CTPAS", "500", 0.0),
        ("IRRF Sobre Folha", "CREDITO", "CTPAS", "501", 0.0),
        ("PIS Folha", "DEBITO", "ADM01", "*056", 0.0),
        ("PIS Folha", "CREDITO", "CTPAS", "501", 0.0),
    ]
    assert {x["titulo"] for x in corpo["linhas"]} == {"Folha Conselho/Diretoria"}
    assert {x["observacoes"] for x in corpo["linhas"]} == {"Folha agosto 26 Conselho"}


def test_item_do_conselho_e_facil_de_mudar(client):
    ev = [e for e in client.get("/eventos", params={"grupo": "conselho", "gera_lancamento": True}).json()
          if e["item_despesa"] == 56]
    assert len(ev) == 3 and ev[0]["observacao"]
    corpo = {k: v for k, v in ev[0].items() if k not in ("id", "ativo", "apelidos")}
    corpo["item_despesa"] = 70
    assert client.put(f"/eventos/{ev[0]['id']}", json=corpo).json()["item_despesa"] == 70
    pdf = Path(__file__).parent / "fixtures" / "folha_socios_08.pdf"
    linhas = client.post("/folha/processar", files={"arquivo": ("s.pdf", pdf.read_bytes(), "application/pdf")}).json()["linhas"]
    assert linhas[0]["codigo_item_despesa"] == "*070"


def _lote(client, rota, nomes, extra=None, centros=None):
    fx = Path(__file__).parent / "fixtures"
    files = [("arquivos", (n, (fx / n).read_bytes(), "application/pdf")) for n in nomes]
    data = {"centros": centros} if centros else {}
    return client.post(rota, files=files, data={**data, **(extra or {})})


NOMES = ["folha_administrativo_08.pdf", "folha_coleta_08.pdf", "folha_oficina_08.pdf", "folha_posto_08.pdf",
         "folha_socios_08.pdf"]


def test_lote_resumo(client):
    r = _lote(client, "/folha/lote/processar", NOMES)
    assert r.status_code == 200
    arq = r.json()["arquivos"]
    assert [a["centro_custo"]["codigo"] for a in arq] == ["ADM01", "ADM02", "ADM04", "ADM03", "ADM01"]
    assert all(a["ok"] and a["conferencias_ok"] and a["eventos_novos"] == [] for a in arq)
    assert arq[4]["perfil"] == "conselho" and arq[4]["nome_csv"] == "lancamentos_Conselho_08-2026.csv"


def test_lote_zip_com_um_csv_por_arquivo(client):
    import io, zipfile
    r = _lote(client, "/folha/lote/csv", NOMES)
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert sorted(z.namelist()) == sorted([
        "lancamentos_Administração_08-2026.csv", "lancamentos_Coleta_08-2026.csv", "lancamentos_Oficina_08-2026.csv",
        "lancamentos_Posto_08-2026.csv", "lancamentos_Conselho_08-2026.csv"])
    linhas = z.read("lancamentos_Coleta_08-2026.csv").decode("utf-8-sig").splitlines()
    assert len(linhas) == 50 and linhas[1].startswith("Folha Coleta;Salário Base;DEBITO;ADM02; 1-1;*070")


def test_lote_arquivo_ruim_nao_derruba_os_outros(client):
    import io, zipfile
    fx = Path(__file__).parent / "fixtures"
    files = [("arquivos", ("ok.pdf", (fx / "folha_coleta_08.pdf").read_bytes(), "application/pdf")),
             ("arquivos", ("ruim.pdf", b"nao e pdf", "application/pdf"))]
    resumo = client.post("/folha/lote/processar", files=files).json()["arquivos"]
    assert [a["ok"] for a in resumo] == [True, False] and resumo[1]["erro"]
    z = zipfile.ZipFile(io.BytesIO(client.post("/folha/lote/csv", files=files).content))
    assert sorted(z.namelist()) == ["LEIAME-erros.txt", "lancamentos_Coleta_08-2026.csv"]
    assert "ruim.pdf" in z.read("LEIAME-erros.txt").decode()


def test_lote_mesmo_centro_duas_vezes_nao_sobrescreve(client):
    import io, zipfile
    r = _lote(client, "/folha/lote/csv", ["folha_coleta_08.pdf", "folha_coleta_08.pdf"])
    assert sorted(zipfile.ZipFile(io.BytesIO(r.content)).namelist()) == [
        "lancamentos_Coleta_08-2026.csv", "lancamentos_Coleta_08-2026_2.csv"]


def test_lote_todos_com_erro(client):
    files = [("arquivos", ("x.pdf", b"nada", "application/pdf"))]
    assert client.post("/folha/lote/csv", files=files).status_code == 422


def test_centro_vem_do_nome_do_arquivo_antes_do_cabecalho(client):
    # PDF da Oficina enviado com nome de arquivo de outro departamento: vale o nome do arquivo, com aviso
    pdf = Path(__file__).parent / "fixtures" / "folha_oficina_08.pdf"
    r = client.post("/folha/processar", files={"arquivo": ("Relatório COOPERNORTE Coleta 08-2026.pdf",
                                                           pdf.read_bytes(), "application/pdf")})
    corpo = r.json()
    assert corpo["centro_custo"]["codigo"] == "ADM02" and corpo["origem_centro"] == "nome do arquivo"
    assert any("Coleta" in a and "OFICINA" in a for a in corpo["avisos"])
    assert {l["codigo_centro_de_custo"] for l in corpo["linhas"]} == {"ADM02", "CTPAS"}


def test_nome_do_arquivo_sem_departamento_usa_o_pdf(client):
    pdf = Path(__file__).parent / "fixtures" / "folha_oficina_08.pdf"
    r = client.post("/folha/processar", files={"arquivo": ("relatorio.pdf", pdf.read_bytes(), "application/pdf")})
    assert r.json()["centro_custo"]["codigo"] == "ADM04" and r.json()["origem_centro"] == "cabeçalho do PDF"


def test_escolha_manual_vence_o_nome_do_arquivo(client):
    pdf = Path(__file__).parent / "fixtures" / "folha_oficina_08.pdf"
    r = client.post("/folha/processar", files={"arquivo": ("Posto.pdf", pdf.read_bytes(), "application/pdf")},
                    data={"centro_custo": "ADM05"})
    assert r.json()["centro_custo"]["codigo"] == "ADM05" and r.json()["origem_centro"] == "escolhido"


@pytest.mark.parametrize("nome,centro", [
    ("Relatório COOPERNORTE Administrativo 08-2026.pdf", "ADM01"), ("Relatório_COOPERNORTE_Coleta_08-2026.pdf", "ADM02"),
    ("Relat#U00f3rio COOPERNORTE Oficina 08-2026.pdf", "ADM04"), ("folha POSTO agosto.pdf", "ADM03"),
    ("Folha Administração.pdf", "ADM01"), ("almoxarifado.pdf", "ADM05")])
def test_centro_pelo_nome_do_arquivo(db, nome, centro):
    from app.folha import sugerir_centro_por_arquivo
    from app.seed.carregar import carregar
    carregar(db)
    assert sugerir_centro_por_arquivo(db, nome).codigo == centro


# ---------- pré-checagem, composição dos totais e recorte do PDF ----------
def lote(client):
    return client.post("/folha/lote/processar", files=[("arquivos", ("folha.pdf", PDF.read_bytes(), "application/pdf"))],
                       data={"centros": "ADM04"}).json()["arquivos"][0]


def test_pre_checagem_ok_sem_problemas(client):
    r = lote(client)
    assert r["ok"] and r["problemas"] == [] and r["linhas"]


def test_composicao_soma_o_total(client):
    r = lote(client)
    for c in r["conferencias"]:
        soma = sum(p["valor"] if p["sinal"] == "+" else -p["valor"] for p in c["composicao"])
        assert round(soma, 2) == c["calculado"]


def test_evento_novo_vira_problema_que_bloqueia(client):
    eid = next(e["id"] for e in client.get("/eventos").json() if e["nome_pdf"] == "Salário Base")
    client.delete(f"/eventos/{eid}")
    r = lote(client)
    assert not r["ok"] and {p["tipo"] for p in r["problemas"]} >= {"evento_novo", "conferencia"}


def test_trecho_devolve_png(client):
    r = client.post("/folha/trecho", files={"arquivo": ("f.pdf", PDF.read_bytes())}, data={"nome": "Salário Base", "origem": "resumo"})
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"
    assert client.post("/folha/trecho", files={"arquivo": ("f.pdf", PDF.read_bytes())}, data={"nome": "Nada", "origem": "resumo"}).status_code == 404


def test_xlsx_mantem_1_1_como_texto(client):
    import io
    from openpyxl import load_workbook
    r = enviar(client, rota="/folha/csv", formato="xlsx")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.xlsx"')
    ws = load_workbook(io.BytesIO(r.content)).active
    cab = [c.value for c in ws[1]]
    linha = {k: ws.cell(2, i + 1) for i, k in enumerate(cab)}
    assert linha["numero_conta"].value == " 1-1" and linha["numero_conta"].number_format == "@"
    assert linha["banco_liquidacao"].value == " 1-1" and linha["banco_liquidacao"].number_format == "@"
    assert linha["codigo_item_despesa"].value.startswith("*0") and linha["data_inicial"].number_format == "DD/MM/YYYY"
    assert isinstance(linha["valor"].value, float)


def test_codigo_do_item_tem_3_digitos():
    from app.folha import codigo_item
    assert [codigo_item(c) for c in (70, 75, 9, 501, 520)] == ["*070", "*075", "*009", "501", "520"]


def test_socios_reconhecido_pelo_departamento_ou_nome():
    import pandas as pd
    from app.folha import detectar_grupo
    df = pd.DataFrame(columns=["Tipo", "Código", "Nome", "Ativos", "Total"])  # layout com coluna Ativos
    assert detectar_grupo(df, "DEPARTAMENTO UNICO") == "conselho"
    assert detectar_grupo(df, "XYZ", "Relatório Folha COOPERNORTE Sócios.pdf") == "conselho"
    assert detectar_grupo(df, "OFICINA", "Relatório Oficina.pdf") == "folha"
