from pathlib import Path

import pytest

from app import config, raiox
from app.routers import raiox as rota

EXTRATO = Path(__file__).parent / "fixtures" / "extrato_junho_2026.xlsx"


@pytest.fixture()
def rx(client, monkeypatch):
    monkeypatch.setattr(config, "RAIOX_SENHA", "secreta123")
    token = client.post("/raiox/login", json={"senha": "secreta123"}).json()["token"]
    return {"X-Raiox-Token": token}


def _arq():
    return {"arquivo": ("Junho_2026.xlsx", EXTRATO.read_bytes())}


def test_senha_do_raiox(client, monkeypatch):
    monkeypatch.setattr(config, "RAIOX_SENHA", "secreta123")
    assert client.post("/raiox/login", json={"senha": "errada"}).status_code == 403
    assert client.get("/raiox/meses").status_code == 403
    assert client.get("/raiox/meses", headers={"X-Raiox-Token": "x"}).status_code == 403
    assert client.post("/raiox/login", json={"senha": " secreta123\n"}).status_code == 200


def test_historico_comeca_com_epoca_boa(client, rx):
    meses = client.get("/raiox/meses", headers=rx).json()
    assert [m["id"] for m in meses] == ["2025-11", "2025-12", "2026-01", "2026-02"]
    assert all(m["formato"] == "antigo" for m in meses)


def test_regras_de_junho():
    doc, own = raiox.analisar(raiox.ler_planilha(EXTRATO.read_bytes()))
    assert doc["id"] == "2026-06"
    assert round(doc["receita"]) == 304613
    assert round(doc["resultado"]) == -111375
    chaves = {f["key"] for f in doc["flags"]}
    assert {"conta-garantida", "posto", "garantida-posto", "duplicados", "falta-quinzena:Interbairros"} <= chaves
    assert len(doc["lanc"]["dup"]) == 21


def test_analisar_compara_com_historico_e_salvar(client, rx):
    prev = client.post("/raiox/analisar", files=_arq(), headers=rx).json()
    assert prev["jaExiste"] is False
    picos = {a["key"] for a in prev["anomalias"]}
    assert "pico:Almoxarifado (estoque matriz)" in picos and "queda-receita" in picos
    assert prev["risco"] == 100
    assert len(client.get("/raiox/meses", headers=rx).json()) == 4  # prévia não salva
    client.post("/raiox/salvar", files=_arq(), headers=rx)
    client.post("/raiox/salvar", files=_arq(), headers=rx)  # mesmo mês substitui
    meses = client.get("/raiox/meses", headers=rx).json()
    assert [m["id"] for m in meses][-1] == "2026-06" and len(meses) == 5
    assert client.delete("/raiox/meses/2026-06", headers=rx).status_code == 200
    assert len(client.get("/raiox/meses", headers=rx).json()) == 4


def test_planilha_errada(client, rx):
    r = client.post("/raiox/analisar", files={"arquivo": ("x.xlsx", b"nao e planilha")}, headers=rx)
    assert r.status_code == 422


def test_parecer(client, rx, monkeypatch):
    client.post("/raiox/salvar", files=_arq(), headers=rx)
    pedidos = []
    monkeypatch.setattr(rota, "pedir_parecer", lambda p: pedidos.append(p) or "RESUMO\nBuraco de R$ 111 mil.")
    r = client.post("/raiox/meses/2026-06/parecer", headers=rx).json()
    assert r["parecer"].startswith("RESUMO") and "jun/26" in pedidos[0]
    junho = client.get("/raiox/meses", headers=rx).json()[-1]
    assert junho["parecer"].startswith("RESUMO")


def test_parecer_sem_chave(client, rx, monkeypatch):
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "")
    client.post("/raiox/salvar", files=_arq(), headers=rx)
    assert client.post("/raiox/meses/2026-06/parecer", headers=rx).status_code == 503


# ---------- análise do período inteiro ----------
def _mes(mid, rot, fin, passagem, rec=None, resultado=0.0):
    return {"id": mid, "rotulo": rot, "formato": "novo", "receita": 300000, "despesas": -300000, "fixas": 250000,
            "passagem": sum(p["liq"] for p in passagem), "resultado": resultado, "financiamento": sum(g["in"] + g["out"] for g in fin),
            "finG": fin, "pass": passagem, "receitaDet": rec or [], "receitaLinhas": {}, "flags": [], "anomalias": []}


def test_completa_compensa_emprestimo_entre_meses():
    jun = _mes("2026-06", "jun/26", [{"g": "Conta do Posto", "in": 0, "out": -100000, "n": 1}],
               [{"n": "FORNECEDORES A PAGAR DDA", "in": 50000, "out": -90000, "liq": -40000}])
    jul = _mes("2026-07", "jul/26", [{"g": "Conta do Posto", "in": 100000, "out": 0, "n": 1}],
               [{"n": "FORNECEDORES A PAGAR DDA", "in": 60000, "out": -20000, "liq": 40000}])
    c = raiox.analise_completa([jun, jul])
    posto = next(x for x in c["contas"] if x["g"] == "Conta do Posto")
    assert posto["aberto"] == 0 and [m["acum"] for m in posto["meses"]] == [-100000, 0]
    assert "completa:posto" not in {f["key"] for f in c["flags"]}
    assert c["passagem"][0]["situacao"] == "compensou"
    assert "completa:compensou" in {f["key"] for f in c["flags"]}


def test_completa_divida_que_nao_volta():
    jun = _mes("2026-06", "jun/26", [{"g": "Conta garantida (crédito do banco)", "in": 300000, "out": 0, "n": 1}],
               [{"n": "REPASSE AO PERMISSIONARIO", "in": 100, "out": -60000, "liq": -59900}])
    jul = _mes("2026-07", "jul/26", [], [{"n": "REPASSE AO PERMISSIONARIO", "in": 0, "out": 0, "liq": 0}])
    c = raiox.analise_completa([jun, jul])
    chaves = {f["key"] for f in c["flags"]}
    assert "completa:garantida" in chaves and "completa:passagem:REPASSE AO PERMISSIONARIO" in chaves
    assert c["dividaEmAberto"] == 300000


def test_completa_quinzena_paga_no_mes_seguinte_nao_falta():
    det = lambda linha, q, lag, v: {"d": "05/07/2026", "linha": linha, "ref": f"{q}ª quinz. x", "lag": lag, "v": v}  # noqa: E731
    jun = _mes("2026-06", "jun/26", [], [], rec=[det(l, q, 1, 1) for l in raiox.LINHAS for q in "12"])  # todo maio
    jul = _mes("2026-07", "jul/26", [], [], rec=[det(l, "1", 1, 1) for l in raiox.LINHAS])  # só 1ª de junho
    c = raiox.analise_completa([jun, jul])
    esperado = {(l, "2ª quinz. jun/26") for l in raiox.LINHAS} | {(l, "1ª quinz. jul/26") for l in raiox.LINHAS[:2]}
    assert {(f["linha"], f["ref"]) for f in c["faltando"]} == esperado


def test_rota_completa(client, rx, monkeypatch):
    assert client.get("/raiox/completa", headers=rx).json()["vazio"] is True
    client.post("/raiox/salvar", files=_arq(), headers=rx)
    c = client.get("/raiox/completa", headers=rx).json()
    assert c["periodo"] == "jun/26 a jun/26" and c["finLanc"]
    assert {"completa:garantida", "completa:posto", "completa:quinzenas"} <= {f["key"] for f in c["flags"]}
    monkeypatch.setattr(rota, "pedir_parecer", lambda p: "TAMANHO REAL DO ROMBO\nX")
    assert client.post("/raiox/completa/parecer", headers=rx).json()["parecerMeses"] == "jun/26 a jun/26"
    assert client.get("/raiox/completa", headers=rx).json()["parecer"].startswith("TAMANHO")


def test_relatorio_excel(client, rx):
    import io
    from openpyxl import load_workbook
    client.post("/raiox/salvar", files=_arq(), headers=rx)
    r = client.get("/raiox/relatorio.xlsx", headers=rx)
    assert r.status_code == 200 and "Relatorio_Tapa_Buraco_" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumo geral", "Jun 2026", "Época boa (ref.)"]
    assert any("=SUMIF(" in str(c.value) for row in wb["Resumo geral"].iter_rows() for c in row)
    assert client.get("/raiox/relatorio.xlsx").status_code == 403
