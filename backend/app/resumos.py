"""Resumos financeiros (Operação Resumos, na área secreta).

Pega o "Relatorio geral" bruto exportado pelo sistema (uma planilha por mês), deixa só as colunas úteis, troca os
valores em texto ("R$ 1.132,11") por número e fecha a conta do mês:

    Sub saldo = total de créditos - total de débitos
    Saldo     = sub saldo - custo do motorista 1 - custo do motorista 2 - permissionário

Os custos fixos (motoristas e permissionário) vêm de faixas de meses guardadas no banco, por exemplo
"abr/25 a mar/26 = 3.350,00". Elas podem ser digitadas, coladas como texto (mensagens do WhatsApp) ou lidas de
uma planilha com linhas do tipo "3350,00 - 01/25".
"""
import io
import re
import unicodedata
import zipfile
from collections import Counter
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CATEGORIAS = ("motorista1", "motorista2", "permissionario")
NOME_CAT = {"motorista1": "Custo Motorista 1", "motorista2": "Custo Motorista 2", "permissionario": "Permissionário"}
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
CURTO = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn").lower().strip()


def rotulo(mid: str) -> str:
    a, m = mid.split("-")
    return f"{CURTO[int(m) - 1]}/{a[2:]}"


def nome_mes(mid: str) -> str:
    a, m = mid.split("-")
    return f"{MESES[int(m) - 1].capitalize()} de {a}"


def valor_br(v) -> float:
    """'R$ 1.132,11' -> 1132.11. Número passa direto; vazio ou '-' vira 0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("R$", "").replace("\xa0", "").replace(" ", "").strip()
    if s in ("", "-"):
        return 0.0
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("-()")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    n = float(s)
    return -n if neg else n


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{2,4})", str(v or ""))
    if not m:
        return None
    d, mes, a = (int(x) for x in m.groups())
    return date(a + 2000 if a < 100 else a, mes, d)


# ------------------------------------------------------------------ leitura do relatório bruto
COLUNAS = {"cap": "cap", "titulo": "titulo", "descricao": "descricao", "data de pagamento": "data",
           "valor credito": "credito", "valor debito": "debito"}


def ler_relatorio(conteudo: bytes, arquivo: str = "") -> dict:
    """Lê o 'Relatorio geral' do sistema. Devolve o mês (o mais frequente nas datas de pagamento) e as linhas."""
    try:
        ws = load_workbook(io.BytesIO(conteudo), data_only=True).worksheets[0]
    except Exception as e:
        raise ValueError(f"{arquivo or 'arquivo'}: não é uma planilha .xlsx válida ({e})")
    linhas_ws = list(ws.iter_rows(values_only=True))
    cab, idx = None, {}
    for i, r in enumerate(linhas_ws[:15]):
        nomes = [_sem_acento(c) if c is not None else "" for c in r]
        if "titulo" in nomes and "data de pagamento" in nomes:
            cab = i
            for j, n in enumerate(nomes):  # se a coluna se repete (planilha já tratada), fica a primeira
                if n in COLUNAS and COLUNAS[n] not in idx:
                    idx[COLUNAS[n]] = j
            break
    if cab is None or not {"titulo", "data", "credito", "debito"} <= idx.keys():
        raise ValueError(f"{arquivo or 'arquivo'}: não achei o cabeçalho do relatório (Titulo, Data de Pagamento, Valor Crédito, Valor Débito).")

    def campo(r, k):
        j = idx.get(k)
        return r[j] if j is not None and j < len(r) else None

    linhas = []
    for r in linhas_ws[cab + 1:]:
        tit, dt = campo(r, "titulo"), _data(campo(r, "data"))
        if tit in (None, "") or dt is None:  # linha de total, vazia ou rodapé
            continue
        linhas.append({"cap": campo(r, "cap"), "titulo": str(tit).strip(), "descricao": campo(r, "descricao"), "data": dt,
                       "credito": round(valor_br(campo(r, "credito")), 2), "debito": round(valor_br(campo(r, "debito")), 2)})
    if not linhas:
        raise ValueError(f"{arquivo or 'arquivo'}: o relatório não tem lançamentos.")
    (a, m), _ = Counter((l["data"].year, l["data"].month) for l in linhas).most_common(1)[0]
    mid = f"{a:04d}-{m:02d}"
    fora = sum(1 for l in linhas if (l["data"].year, l["data"].month) != (a, m))
    return {"id": mid, "rotulo": rotulo(mid), "nome": nome_mes(mid), "arquivo": arquivo, "linhas": linhas,
            "credito": round(sum(l["credito"] for l in linhas), 2), "debito": round(sum(l["debito"] for l in linhas), 2),
            "foraDoMes": fora}


# ------------------------------------------------------------------ custos fixos por faixa de meses
def custo_do_mes(faixas: list[dict], mid: str) -> float | None:
    """Valor da faixa que cobre o mês. Se mais de uma cobre, vale a última cadastrada."""
    for f in reversed(faixas or []):
        if f["de"] <= mid and (not f.get("ate") or mid <= f["ate"]):
            return float(f["valor"])
    return None


def calcular(rel: dict, custos: dict) -> dict:
    sub = round(rel["credito"] - rel["debito"], 2)
    vals = {c: custo_do_mes(custos.get(c, []), rel["id"]) for c in CATEGORIAS}
    faltando = [NOME_CAT[c] for c, v in vals.items() if v is None]
    saldo = round(sub - sum(v or 0 for v in vals.values()), 2)
    avisos = []
    if faltando:
        avisos.append(f"Sem valor cadastrado para {', '.join(faltando)} em {rel['rotulo']} (contado como 0).")
    if rel.get("foraDoMes"):
        avisos.append(f"{rel['foraDoMes']} lançamento(s) com data fora de {rel['rotulo']}.")
    return {"id": rel["id"], "rotulo": rel["rotulo"], "nome": rel["nome"], "arquivo": rel["arquivo"], "lancamentos": len(rel["linhas"]),
            "credito": rel["credito"], "debito": rel["debito"], "subSaldo": sub, **vals, "saldo": saldo, "avisos": avisos}


# ------------------------------------------------------------------ leitura de custos em texto livre / planilha
_NOMES_MES = {}
for _i, (_n, _c) in enumerate(zip(MESES, CURTO), 1):
    _NOMES_MES[_sem_acento(_n)] = _i
    _NOMES_MES[_c] = _i
_RE_MES = "|".join(sorted(_NOMES_MES, key=len, reverse=True))
_RE_VALOR = re.compile(r"(?:r\$\s*)?(\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?|\d+,\d{1,2}|\d{3,}(?:\.\d{1,2})?)(?![\d/])")
_RE_TOKEN = re.compile(rf"(?P<mm>\b\d{{1,2}})/(?P<aa>\d{{4}}|\d{{2}})\b|\b(?P<nome>{_RE_MES})\b(?:\s*(?:de|/)?\s*(?P<ano>\d{{4}}|\d{{2}})\b)?")


def categorias_da_linha(s: str) -> list[str] | None:
    t = _sem_acento(s)
    if "permission" in t:
        return ["permissionario"]
    if "motorista" in t:
        if re.search(r"motorista\s*(1|um|01)\b", t):
            return ["motorista1"]
        if re.search(r"motorista\s*(2|dois|02)\b", t):
            return ["motorista2"]
        return ["motorista1", "motorista2"]
    return None


def _ano(a: str | None) -> int | None:
    if not a:
        return None
    n = int(a)
    return n + 2000 if n < 100 else n


def ler_linha(s: str) -> dict | None:
    """'Abril 2025 a março 2026 3.350,00' -> {datas: [(4,2025),(3,2026)], modo: 'faixa', valor: 3350.0}.
    Também entende '3350,00 - 01/25', '01/2025 a 03/2025 3050', 'a partir de outubro 2026 4.100,00'."""
    t = _sem_acento(s)
    t = re.sub(r"\b\d{1,2}:\d{2}\b", " ", t).replace("encaminhada", " ")  # hora e marcação do WhatsApp
    datas, modo, resto = [], "faixa", t
    for m in _RE_TOKEN.finditer(t):
        if m.group("mm"):
            mes, ano = int(m.group("mm")), _ano(m.group("aa"))
            if not 1 <= mes <= 12:
                continue
        else:
            mes, ano = _NOMES_MES[m.group("nome")], _ano(m.group("ano"))
        datas.append([mes, ano, m.start(), m.end()])
    if not datas:
        return None
    for d in reversed(datas):  # tira as datas antes de procurar o valor
        resto = resto[:d[2]] + " " * (d[3] - d[2]) + resto[d[3]:]
    vals = _RE_VALOR.findall(resto)
    if not vals:
        return None
    valor = valor_br(vals[-1])
    if len(datas) >= 2:
        entre = t[datas[0][3]:datas[1][2]]
        if re.search(r"\be\b|,", entre) and not re.search(r"\b(a|ate|à)\b|-", entre):
            modo = "lista"
    if re.search(r"a partir|desde|em diante|pra frente|para frente", t):
        modo = "aberto"
    return {"datas": [(d[0], d[1]) for d in datas], "modo": modo, "valor": round(valor, 2)}


def _completar_anos(itens: list[dict]) -> None:
    """Preenche os anos que faltam ('Janeiro a março 3.050,00') olhando a própria faixa e as vizinhas."""
    hoje = date.today().year
    for it in itens:
        ds = it["datas"]
        for i in range(len(ds)):  # dentro da mesma linha: herda do vizinho mais próximo com ano
            if ds[i][1] is None:
                for j in list(range(i + 1, len(ds))) + list(range(i - 1, -1, -1)):
                    if ds[j][1] is not None:
                        y = ds[j][1]
                        y = y if (j > i and ds[i][0] <= ds[j][0]) or (j < i and ds[i][0] >= ds[j][0]) else (y - 1 if j > i else y + 1)
                        ds[i] = (ds[i][0], y)
                        break
    for k, it in enumerate(itens):  # linha inteira sem ano: usa a linha seguinte (ou a anterior)
        if all(d[1] is None for d in it["datas"]):
            ds = it["datas"]
            prox = next((x["datas"][0] for x in itens[k + 1:] if x["datas"][0][1] is not None), None)
            ant = next((x["datas"][-1] for x in reversed(itens[:k]) if x["datas"][-1][1] is not None), None)
            if prox:
                fim = prox[1] if ds[-1][0] < prox[0] else prox[1] - 1
            elif ant:
                fim = ant[1] if ds[-1][0] > ant[0] else ant[1] + 1
            else:
                fim = hoje
            novas, y = [], fim
            for i in range(len(ds) - 1, -1, -1):
                if i < len(ds) - 1 and ds[i][0] > ds[i + 1][0]:
                    y -= 1
                novas.insert(0, (ds[i][0], y))
            it["datas"] = novas


def _mid(d) -> str:
    return f"{d[1]:04d}-{d[0]:02d}"


def _para_faixas(it: dict) -> list[dict]:
    ds, v = it["datas"], it["valor"]
    if it["modo"] == "aberto":
        return [{"de": _mid(ds[0]), "ate": None, "valor": v}]
    if it["modo"] == "lista" or len(ds) == 1:
        return [{"de": _mid(d), "ate": _mid(d), "valor": v} for d in ds]
    de, ate = sorted([_mid(ds[0]), _mid(ds[-1])])
    return [{"de": de, "ate": ate, "valor": v}]


def ler_texto(texto: str, categoria: str = "auto") -> dict:
    """Texto livre (uma faixa por linha). Linhas como 'Permissionário' ou 'Motorista 1' mudam a categoria das seguintes.
    categoria: 'auto' (só pelo texto), uma das CATEGORIAS ou 'motoristas' (os dois)."""
    padrao = ["motorista1", "motorista2"] if categoria == "motoristas" else ([categoria] if categoria in CATEGORIAS else None)
    atual, itens, ignoradas, sem_cat = padrao, [], [], []
    for linha in (texto or "").splitlines():
        if not linha.strip():
            continue
        lido = ler_linha(linha)
        cats = categorias_da_linha(linha)
        if lido is None:
            if cats and categoria == "auto":
                atual = cats
            elif _sem_acento(linha) not in ("encaminhada", "") and not re.fullmatch(r"\s*\d{1,2}:\d{2}\s*", linha):
                ignoradas.append(linha.strip())
            continue
        usar = (cats if categoria == "auto" and cats else None) or atual
        if not usar:
            sem_cat.append(linha.strip())
            continue
        lido["cats"] = usar
        itens.append(lido)
    # anos são completados por categoria, na ordem em que aparecem
    for c in set(c for it in itens for c in it["cats"]):
        _completar_anos([it for it in itens if c in it["cats"]])
    faixas = {c: [] for c in CATEGORIAS}
    for it in itens:
        for c in it["cats"]:
            faixas[c].extend(_para_faixas(it))
    return {"faixas": faixas, "ignoradas": ignoradas, "semCategoria": sem_cat}


def _texto_celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return f"{v.month:02d}/{v.year}"
    if isinstance(v, float):
        return f"{v:.2f}".replace(".", ",")
    return str(v)


def ler_planilha_custos(conteudo: bytes, categoria: str = "auto") -> dict:
    """Planilha com os valores mês a mês ('3350,00 - 01/25' numa célula, ou valor e mês em células vizinhas).
    A categoria sai do nome da aba ou do cabeçalho da coluna (Permissionário, Motorista 1...), ou da escolha na tela."""
    wb = load_workbook(io.BytesIO(conteudo), data_only=True)
    linhas = []
    for ws in wb.worksheets:
        cat_aba = categorias_da_linha(ws.title) if categoria == "auto" else None
        cat_col: dict[int, list[str]] = {}
        atual_aba = cat_aba
        for r in ws.iter_rows(values_only=True):
            cels = [_texto_celula(v) for v in r]
            usou = False
            for j, s in enumerate(cels):
                if not s.strip():
                    continue
                if ler_linha(s) is not None:
                    linhas.append((cat_col.get(j) or atual_aba, s))
                    usou = True
                elif categoria == "auto" and categorias_da_linha(s):
                    cat_col[j] = categorias_da_linha(s)
                    if len([x for x in cels if x.strip()]) == 1:
                        atual_aba = cat_col[j]
            junto = " ".join(x for x in cels if x.strip())
            if not usou and junto and ler_linha(junto) is not None:
                j0 = next(j for j, x in enumerate(cels) if x.strip())
                linhas.append((cat_col.get(j0) or atual_aba, junto))
    rotulos = {("permissionario",): "Permissionário", ("motorista1",): "Motorista 1", ("motorista2",): "Motorista 2",
               ("motorista1", "motorista2"): "Motoristas"}
    texto, cat_ant = [], None
    for cats, s in linhas:  # reaproveita o leitor de texto: um rótulo de categoria antes de cada troca
        if categoria == "auto" and cats and cats != cat_ant:
            texto.append(rotulos[tuple(cats)])
            cat_ant = cats
        texto.append(s)
    return ler_texto("\n".join(texto), categoria)


# ------------------------------------------------------------------ planilhas de saída
AZUL, VERMELHO = "1F5FA8", "C00000"
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
FMT = '#,##0.00;"- "#,##0.00;"-"'
CAB_FILL = PatternFill("solid", fgColor="1F3864")


def _salvar(wb: Workbook) -> bytes:
    wb.calculation.fullCalcOnLoad = True  # o Excel recalcula as fórmulas ao abrir
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _bloco_contas(ws, linha: int, col_rot: int, col_val: int, cred_ref: str, deb_ref: str, r: dict) -> int:
    """Escreve Sub Saldo, custos e Saldo (como no relatório manual). Devolve a linha seguinte."""
    def put(rot, val, cor_rot, cor_val, negrito_val=True):
        nonlocal linha
        a = ws.cell(linha, col_rot, rot)
        a.font = Font(bold=True, color=cor_rot)
        b = ws.cell(linha, col_val, val)
        b.font, b.number_format = Font(bold=negrito_val, color=cor_val), FMT
        linha += 1
        return b.coordinate
    sub = put("Sub Saldo:", f"={cred_ref}-{deb_ref}", AZUL, VERMELHO)
    refs = [put(NOME_CAT[c], -(r[c] or 0), "000000", AZUL) for c in CATEGORIAS]
    put("Saldo:", f"=SUM({sub}:{refs[-1]})", AZUL, VERMELHO)
    return linha


def xlsx_mes(rel: dict, r: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = rel["rotulo"].replace("/", "-")
    a, m = map(int, rel["id"].split("-"))
    fim = (date(a + (m == 12), m % 12 + 1, 1) - date.resolution)
    ws["A1"] = f"Relatorio geral de 01/{m:02d}/{a % 100:02d} a {fim.day:02d}/{m:02d}/{a % 100:02d}"
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:F1")
    cab = ["Cap", "Titulo", "Descrição", "Data de Pagamento", "Valor Crédito", "Valor Débito"]
    for j, t in enumerate(cab, 1):
        c = ws.cell(2, j, t)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), CAB_FILL, BORDA
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, l in enumerate(rel["linhas"], 3):
        vals = [l["cap"], l["titulo"], l["descricao"], l["data"], l["credito"], l["debito"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.border = BORDA
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 3))
        ws.cell(i, 4).number_format = "DD/MM/YYYY"
        ws.cell(i, 5).number_format = ws.cell(i, 6).number_format = FMT
    ult = 2 + len(rel["linhas"])
    t = ult + 1
    ws.cell(t, 4, "Total:").font = Font(bold=True)
    for j, col in ((5, "E"), (6, "F")):
        c = ws.cell(t, j, f"=SUM({col}3:{col}{ult})")
        c.font, c.number_format, c.border = Font(bold=True, color=VERMELHO, size=12), FMT, BORDA
    _bloco_contas(ws, t + 1, 5, 6, f"E{t}", f"F{t}", r)
    for col, w in zip("ABCDEF", (7, 42, 52, 18, 16, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:F{ult}"
    ws.page_setup.orientation = "landscape"  # impressão: cabe na largura da folha
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "2:2"
    return _salvar(wb)


def xlsx_resumo(lista: list[dict]) -> bytes:
    lista = sorted(lista, key=lambda r: r["id"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Resumo financeiro · {lista[0]['rotulo']} a {lista[-1]['rotulo']}" if len(lista) > 1 else f"Resumo financeiro · {lista[0]['nome']}"
    ws["A1"].font = Font(bold=True, size=13)
    linha, saldos = 3, []
    for r in lista:
        c = ws.cell(linha, 1, r["nome"])
        c.font, c.fill = Font(bold=True, color="FFFFFF"), CAB_FILL
        ws.cell(linha, 2).fill = ws.cell(linha, 3).fill = CAB_FILL
        ws.cell(linha + 1, 1, "Créditos / Débitos").font = Font(bold=True)
        for j, k in ((2, "credito"), (3, "debito")):
            x = ws.cell(linha + 1, j, r[k])
            x.font, x.number_format = Font(bold=True, color=VERMELHO), FMT
        fim = _bloco_contas(ws, linha + 2, 1, 3, f"B{linha + 1}", f"C{linha + 1}", r)
        saldos.append(f"C{fim - 1}")
        linha = fim + 1
    if len(lista) > 1:
        ws.cell(linha, 1, "Saldo acumulado do período:").font = Font(bold=True, color=AZUL, size=12)
        c = ws.cell(linha, 3, "=" + "+".join(saldos))
        c.font, c.number_format = Font(bold=True, color=VERMELHO, size=12), FMT
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 30, 16, 16

    t = wb.create_sheet("Tabela")  # um mês por linha, bom para filtrar ou fazer gráfico
    cab = ["Mês", "Créditos", "Débitos", "Sub Saldo", "Custo Motorista 1", "Custo Motorista 2", "Permissionário", "Saldo"]
    for j, x in enumerate(cab, 1):
        c = t.cell(1, j, x)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), CAB_FILL, BORDA
    for i, r in enumerate(lista, 2):
        vals = [r["rotulo"], r["credito"], r["debito"], f"=B{i}-C{i}", -(r["motorista1"] or 0), -(r["motorista2"] or 0),
                -(r["permissionario"] or 0), f"=SUM(D{i}:G{i})"]
        for j, v in enumerate(vals, 1):
            c = t.cell(i, j, v)
            c.border = BORDA
            if j > 1:
                c.number_format = FMT
    n = len(lista) + 2
    t.cell(n, 1, "Total").font = Font(bold=True)
    for j in range(2, 9):
        col = get_column_letter(j)
        c = t.cell(n, j, f"=SUM({col}2:{col}{n - 1})")
        c.font, c.number_format, c.border = Font(bold=True, color=VERMELHO if j == 8 else "000000"), FMT, BORDA
    for j, w in enumerate((10, 15, 15, 15, 18, 18, 16, 15), 1):
        t.column_dimensions[get_column_letter(j)].width = w
    return _salvar(wb)


def nome_arquivo(mid: str) -> str:
    a, m = mid.split("-")
    return f"{MESES[int(m) - 1].replace('ç', 'c')}_{a[2:]}.xlsx"


def gerar_zip(arquivos: list[tuple[str, bytes]], custos: dict) -> tuple[bytes, str]:
    rels = [ler_relatorio(c, n) for n, c in arquivos]
    res = [calcular(r, custos) for r in rels]
    b, usados = io.BytesIO(), Counter()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for rel, r in zip(rels, res):
            nome = nome_arquivo(rel["id"])
            usados[nome] += 1
            if usados[nome] > 1:
                nome = nome.replace(".xlsx", f"_{usados[nome]}.xlsx")
            z.writestr(nome, xlsx_mes(rel, r))
        ids = sorted(r["id"] for r in res)
        sufixo = rotulo(ids[0]).replace("/", "") + (f"_a_{rotulo(ids[-1]).replace('/', '')}" if ids[-1] != ids[0] else "")
        z.writestr(f"resumo_{sufixo}.xlsx", xlsx_resumo(res))
    return b.getvalue(), f"relatorios_{sufixo}.zip"
