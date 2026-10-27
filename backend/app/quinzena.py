"""Quinzena dos permissionários: lê o relatório de remuneração (uma coluna por evento, uma linha por carro) e monta
o lançamento por item, como a aba 'lancamentos' da planilha da quinzena.

Regras (aba Resumo da planilha):
- cada evento vai para uma linha do lançamento pela primeira regra cujo padrão aparece no nome (sem o código);
- Crédito/Pagto quinzenal = saldo negativo da quinzena anterior (evento 342) − soma dos líquidos + adiantamentos;
- Taxa administrativa (pode vir em mais de uma coluna) = Cooperativa + Sindicato, divididas pelos valores por carro;
- Débitos de acordo saem como DÉBITO e voltam como 'Estorno' (CRÉDITO) 15 dias depois;
- o valor do lançamento é sempre positivo (a planilha usa o valor absoluto)."""
import csv
import json
import io
import re
import unicodedata
from datetime import date, timedelta

import pandas as pd
from sqlalchemy import select
from sqlalchemy.orm import Session

from .folha import codigo_item
from .models import Configuracao, LinhaQuinzena, RegraQuinzena

COLUNAS = ["titulo", "descricao", "tipo_lancamento", "codigo_centro_de_custo", "numero_conta", "codigo_item_despesa",
           "codigo_fonte_recurso", "id_forma_pagamento", "num_parcelas", "data_inicial", "dias_intervalo", "valor",
           "observacoes"]
TEXTO_XLSX = ("codigo_centro_de_custo", "numero_conta", "codigo_item_despesa", "codigo_fonte_recurso")
ESPECIAIS = {"_saldo": "Saldo negativo da quinzena anterior (entra no crédito)",
             "_taxa": "Taxa administrativa (dividida em Cooperativa e Sindicato)",
             "_ignorar": "Não entra no lançamento (compõe o líquido)",
             "_terminais": "Débito/crédito das intermunicipais (vai para o terminal do carro, pelo prefixo)"}
GRUPOS = {"INTERBAIRRO": "Interbairros", "ALIMENTADOR": "Alimentadora", "INTERMUNICIPA": "Intermunicipal",
          "TIUMA": "Intermunicipal"}
MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro",
         "Novembro", "Dezembro"]


def norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "").upper())
    return re.sub(r"\s+", " ", "".join(c for c in s if unicodedata.category(c) != "Mn")).strip()


def sem_codigo(nome) -> str:
    """'41862 - BANDEIRA ' -> 'BANDEIRA'"""
    return re.sub(r"^\s*\d+\s*-\s*", "", str(nome)).strip()


def config(db: Session) -> dict:
    return {c.chave: c.valor for c in db.scalars(select(Configuracao))}


# ---------- leitura ----------
def ler_relatorio(conteudo: bytes) -> pd.DataFrame:
    """Acha, em qualquer aba, a tabela com 'Prefixo' e 'Liquido' no cabeçalho; para na linha 'Total'."""
    abas = pd.read_excel(io.BytesIO(conteudo), sheet_name=None, header=None)
    for bruto in abas.values():
        for i in range(min(15, len(bruto))):
            cab = [norm(v) for v in bruto.iloc[i].tolist()]
            if "PREFIXO" in cab and "LIQUIDO" in cab:
                df = bruto.iloc[i + 1:].copy()
                df.columns = [str(v).strip() if not pd.isna(v) else f"_vazia{j}" for j, v in enumerate(bruto.iloc[i])]
                col_pref = df.columns[cab.index("PREFIXO")]
                df = df.rename(columns={col_pref: "Prefixo", df.columns[cab.index("LIQUIDO")]: "Liquido"})
                fim = df.index[df["Prefixo"].astype(str).str.strip().str.upper().isin(["TOTAL", "TOTAIS"])]
                if len(fim):
                    df = df.loc[:fim[0] - 1]
                df = df[df["Prefixo"].notna() & (df["Prefixo"].astype(str).str.strip() != "")]
                return df.loc[:, [c for c in df.columns if not c.startswith("_vazia")]].reset_index(drop=True)
    raise ValueError("Não achei a tabela do relatório (cabeçalho com 'Prefixo' e 'Liquido').")


def eventos_do_relatorio(df: pd.DataFrame) -> list[dict]:
    """Uma entrada por coluna de evento: total e o valor de cada carro (só os diferentes de zero)."""
    fixas = {"PREFIXO", "PRE. HISTORICO", "ANO", "LIQUIDO"}
    saida = []
    for col in df.columns:
        if norm(col) in fixas:
            continue
        valores = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
        por_carro = [{"prefixo": str(p).strip(), "valor": round(float(v), 2)}
                     for p, v in zip(df["Prefixo"], valores) if abs(v) >= 0.005]
        saida.append({"nome": col, "total": round(float(valores.sum()), 2), "por_carro": por_carro})
    return saida


def detectar_grupo(nome_arquivo: str | None, eventos: list[dict]) -> str | None:
    for texto in [nome_arquivo or "", *[e["nome"] for e in eventos if "REMUNERA" in norm(e["nome"])]]:
        t = norm(texto)
        for chave, grupo in GRUPOS.items():
            if chave in t:
                return grupo
    return None


# ---------- classificação ----------
def regra_do_evento(regras, nome: str):
    n = norm(sem_codigo(nome))
    return next((r for r in regras if norm(r.padrao) in n), None)


def processar(db: Session, df: pd.DataFrame, grupo: str, data_lanc: date, observacao: str, conta: str | None = None,
              credito_manual: float | None = None, incluir_zeradas: bool = False) -> dict:
    cfg = config(db)
    t_coop, t_sind = float(cfg.get("taxa_cooperativa", 0)), float(cfg.get("taxa_sindicato", 0))
    conta = conta or cfg.get("conta_quinzena", "")
    regras = db.scalars(select(RegraQuinzena).order_by(RegraQuinzena.ordem, RegraQuinzena.id)).all()
    linhas_cad = db.scalars(select(LinhaQuinzena).order_by(LinhaQuinzena.ordem)).all()
    eventos = eventos_do_relatorio(df)
    liquido = round(float(pd.to_numeric(df["Liquido"], errors="coerce").fillna(0).sum()), 2)

    por_linha: dict[str, list] = {}
    sem_regra, saldo, taxa_eventos, terminais_ev = [], [], [], []
    for ev in eventos:
        r = regra_do_evento(regras, ev["nome"])
        ev["linha"] = r.linha if r else None
        ev["regra"] = r.padrao if r else None
        if r is None:
            if ev["total"]:
                sem_regra.append(ev)
            continue
        if r.linha == "_saldo":
            saldo.append(ev)
        elif r.linha == "_taxa":
            taxa_eventos.append(ev)
        elif r.linha == "_terminais":
            terminais_ev.append(ev)
        elif r.linha != "_ignorar":
            por_linha.setdefault(r.linha, []).append(ev)

    def parcela(nome, valor, ev=None):
        return {"nome": nome, "valor": round(valor, 2), "por_carro": ev["por_carro"] if ev else []}

    composicao: dict[str, list] = {c: [parcela(e["nome"], e["total"], e) for e in evs if e["total"]]
                                   for c, evs in por_linha.items()}
    # débito/crédito das intermunicipais: cada carro vai para o terminal do seu prefixo
    terminais = json.loads(cfg.get("terminais") or "{}")
    sem_terminal = []
    for ev in terminais_ev:
        partes: dict[str, list] = {}
        for c in ev["por_carro"]:
            linha = next((t for t, pref in terminais.items()
                          if any(c["prefixo"].lstrip("0").startswith(p.lstrip("0")) or c["prefixo"].startswith(p) for p in pref)), None)
            if linha is None:
                sem_terminal.append((ev["nome"], c))
            else:
                partes.setdefault(linha, []).append(c)
        for linha, carros_t in partes.items():
            composicao.setdefault(linha, []).append(
                {"nome": ev["nome"], "valor": round(sum(c["valor"] for c in carros_t), 2), "por_carro": carros_t})
    # crédito quinzenal
    soma_saldo = round(sum(e["total"] for e in saldo), 2)
    cred = [parcela("Saldo negativo da quinzena anterior", soma_saldo, saldo[0] if len(saldo) == 1 else None),
            parcela("Menos a soma dos líquidos", -liquido)] + composicao.get("credito", [])
    credito_calc = round(sum(p["valor"] for p in cred), 2)
    if credito_manual is not None:
        cred.append(parcela("Ajuste informado por você", round(credito_manual - credito_calc, 2)))
    composicao["credito"] = cred
    # taxa: carros que pagaram (em qualquer coluna de taxa) e divisão Cooperativa / Sindicato
    por_carro: dict[str, float] = {}
    for e in taxa_eventos:
        for c in e["por_carro"]:
            por_carro[c["prefixo"]] = por_carro.get(c["prefixo"], 0) + c["valor"]
    taxa_total = round(sum(por_carro.values()), 2)
    carros = sum(1 for v in por_carro.values() if abs(v) >= 0.005)
    # cada carro que pagou a taxa paga a parte da Cooperativa; o resto é do Sindicato
    sinal = -1 if taxa_total < 0 else 1
    coop = round(sinal * carros * t_coop, 2)
    composicao["taxa_coop"] = [parcela(f"{carros} carros pagaram a taxa × {t_coop:.2f} (Cooperativa)", coop)] \
        + composicao.get("taxa_coop", [])
    composicao["taxa_sind"] = [parcela("Taxa administrativa paga", taxa_total),
                               parcela("Menos a parte da Cooperativa", -coop)] + composicao.get("taxa_sind", [])
    avisos = []
    esperado = round(carros * (t_coop + t_sind), 2)
    if taxa_eventos and abs(abs(taxa_total) - esperado) > 0.05:
        avisos.append(f"A taxa administrativa somou {abs(taxa_total):.2f}, mas {carros} carros × {t_coop + t_sind:.2f} "
                      f"dariam {esperado:.2f}. A Cooperativa ficou com {carros} × {t_coop:.2f} e o Sindicato com o restante.")

    linhas = []
    for l in linhas_cad:
        comp = composicao.get("acordo" if l.codigo == "estorno_acordo" else l.codigo, [])
        total = round(sum(p["valor"] for p in comp), 2)
        if l.codigo == "credito" and credito_manual is not None:
            total = round(credito_manual, 2)
        if abs(total) < 0.005 and not incluir_zeradas:
            continue
        dia = data_lanc + timedelta(days=15) if l.codigo == "estorno_acordo" else data_lanc
        linhas.append({
            "codigo": l.codigo, "composicao": comp, "total_eventos": total, "dias_depois": 15 if l.codigo == "estorno_acordo" else 0,
            "titulo": f"Deposito Permissionários {grupo}", "descricao": f"{l.descricao} - {grupo}",
            "tipo_lancamento": l.tipo_lancamento, "codigo_centro_de_custo": l.centro, "numero_conta": conta,
            "codigo_item_despesa": codigo_item(l.item) if l.item is not None else None,
            "codigo_fonte_recurso": l.fonte, "id_forma_pagamento": None, "num_parcelas": None,
            "data_inicial": dia, "dias_intervalo": None, "valor": round(abs(total), 2), "observacoes": observacao})
    # o que deve bater com o depósito no banco (como o 'Check-in' da aba Resumo): soma dos valores, sem o estorno
    total_quinzena = round(sum(l["valor"] for l in linhas if l["codigo"] != "estorno_acordo"), 2)
    problemas = [{"tipo": "evento_sem_regra", "titulo": f"Evento sem regra: {sem_codigo(e['nome'])}",
                  "detalhe": f"O evento “{e['nome'].strip()}” somou {e['total']:.2f} em {len(e['por_carro'])} carro(s) e "
                             "nenhuma regra diz em qual linha do lançamento ele entra. Escolha a linha abaixo.",
                  "evento": {"nome": e["nome"].strip(), "sugestao": sem_codigo(e["nome"]).upper(), "total": e["total"],
                             "por_carro": e["por_carro"]}} for e in sem_regra]
    problemas += [{"tipo": "carro_sem_terminal", "titulo": f"Carro {c['prefixo']} sem terminal",
                   "detalhe": f"O carro {c['prefixo']} tem {c['valor']:.2f} em “{nome.strip()}”, mas o prefixo dele não é de "
                              "nenhum terminal. Ajuste os prefixos dos terminais em “Linhas do lançamento e taxa”."}
                  for nome, c in sem_terminal]
    return {"grupo": grupo, "liquido": liquido, "saldo_negativo": soma_saldo, "credito_calculado": credito_calc,
            "taxa": {"total": abs(taxa_total), "carros": carros, "cooperativa": abs(coop),
                     "sindicato": abs(round(taxa_total - coop, 2))},
            "total_quinzena": total_quinzena, "n_carros": len(df), "avisos": avisos, "problemas": problemas,
            "eventos": [{k: e[k] for k in ("nome", "total", "linha", "regra")} for e in eventos],
            "linhas": linhas}


def observacao_padrao(quinzena: int, mes: int, ano: int) -> str:
    return f"{quinzena}° Quinzena de {MESES[mes - 1]} {ano}"


# ---------- saída ----------
def para_xlsx(linhas: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "lancamentos"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for l in linhas:
        ws.append([l[c] for c in COLUNAS])
    for j, nome in enumerate(COLUNAS, start=1):
        for (cel,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
            if nome in TEXTO_XLSX:
                cel.number_format = "@"
                cel.value = None if cel.value is None else str(cel.value)
            elif nome == "data_inicial":
                cel.number_format = "DD/MM/YYYY"
            elif nome == "valor":
                cel.number_format = "#,##0.00"
        ws.column_dimensions[ws.cell(1, j).column_letter].width = max(12, min(55, len(nome) + 4))
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def para_csv(linhas: list[dict]) -> str:
    def fmt(v):
        if v is None:
            return ""
        if isinstance(v, date):
            return v.strftime("%d/%m/%Y")
        if isinstance(v, float):
            return f"{v:.2f}".replace(".", ",")
        return v
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(COLUNAS)
    for l in linhas:
        w.writerow([fmt(l[c]) for c in COLUNAS])
    return buf.getvalue()
