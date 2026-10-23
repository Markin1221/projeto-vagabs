"""Gera os arquivos do seed (app/seed/dados/*.json) a partir da planilha base.

Só precisa rodar de novo quando a planilha base mudar:
    python scripts/gerar_seed_da_planilha.py [caminho_da_planilha.xlsx]

Lê: aba "Contabil" (eventos, centro de custo e item de despesa), aba "lancamentos CSV" (ordem e
descrição das linhas do CSV) e aba "Itens Despesas Resumido" (centros de custo e itens).
Os JSONs gerados são versionados; a planilha (.xlsx) não precisa ir para o repositório.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "app" / "seed" / "dados"
PADRAO = RAIZ.parent / "docs" / "referencia" / "Contabil_folha_Administração.xlsx"

# Linhas da aba Contabil por seção (a planilha não marca isso; é a posição dos blocos)
ADICIONAIS = range(5, 28)
DESCONTOS = range(29, 44)
TOTAIS = (28, 44, 45, 46, 47)
GPS = range(53, 70)
PIS = (72, 73)
FGTS = (85, 86, 88, 89, 90, 93, 94, 95, 96, 97)

# Linhas do CSV que não vêm direto de uma linha do PDF: valor calculado a partir de outros eventos.
# {nome} = valor do evento no PDF (coluna Total); {nome|Coluna} = outra coluna do PDF.
CALCULOS = {
    "TOTAL LÍQUIDO A PAGAR SEM FÉRIAS": "{TOTAL LÍQUIDO A PAGAR} - {Adiantamento} - {TOTAL LÍQUIDO A PAGAR (FÉRIAS)}",
    "Total Líquido INSS - Oficina": (
        "{Sócios} + {Autônomos} + {Empresa Funcionários} + {RAT Emp (RAT x FAP = 5,0721 %)} + {Vlr. Ocorrência}"
        " + {Empresa Sócios} + {Empresa Autônomos} + {Cooperativas} + {Resíduo Mês Anterior} + {Valor Retido}"
        " + {Terceiros 5,80 %}"),
    "PIS Folha - Oficina": "{PIS Folha}",
    "Total FGTS  MÊS": "{Total FGTS apurado recibos s/CS}",
    "Base de calc. Multa FGTS GRRF ": "{Base de calc. Multa FGTS GRRF}",
    "Desconto Empréstimo Crédito Trabalhador": "{Empréstimo (Crédito Trabalhador)|Total} - {Empréstimo (Crédito Trabalhador)|Ativos}",
    "Plano de Saúde - Total ": "{Plano De Saúde - Mens.} * 2",
}


def limpar_nome(v):
    return re.sub(r"\s*:\s*$", "", str(v).strip())


def secao_da_linha(r):
    if r in ADICIONAIS: return "adicional"
    if r in DESCONTOS: return "desconto"
    if r in TOTAIS: return "total"
    if r in GPS: return "gps"
    if r in PIS: return "pis"
    if r in FGTS: return "fgts"
    return "calculado"


def inteiro(v):
    try:
        return int(str(v).lstrip("*").lstrip("0") or 0) if str(v).startswith("*") else int(float(v))
    except (TypeError, ValueError):
        return None


def main(caminho):
    wf = openpyxl.load_workbook(caminho)                 # fórmulas (de onde cada linha do CSV vem)
    wv = openpyxl.load_workbook(caminho, data_only=True)  # valores já calculados pelo Excel
    ct, ctf = wv["Contabil"], wf["Contabil"]
    csv_v, csv_f = wv["lancamentos CSV"], wf["lancamentos CSV"]
    itens_ws = wv["Itens Despesas Resumido"]

    # ---- centros de custo e itens de despesa ----
    centros = []
    for r in range(2, 60):
        cod, nome = itens_ws.cell(r, 7).value, itens_ws.cell(r, 8).value
        if cod and nome:
            tipo = "departamento" if str(cod).startswith("ADM") else "passagem" if cod == "CTPAS" else "receita"
            centros.append({"codigo": str(cod), "nome": str(nome), "tipo": tipo})
    itens = []
    for r in range(2, 200):
        cod, nome = itens_ws.cell(r, 1).value, itens_ws.cell(r, 2).value
        if inteiro(cod) and nome:
            itens.append({"codigo": inteiro(cod), "nome": str(nome).strip(), "grupo": "despesa"})
    for r in range(2, 60):
        cod, nome = itens_ws.cell(r, 10).value, itens_ws.cell(r, 11).value
        if isinstance(cod, (int, float)) and cod >= 500 and nome:
            itens.append({"codigo": int(cod), "nome": str(nome).strip(), "grupo": "passagem"})

    # ---- eventos que viram linha do CSV (na ordem do CSV) ----
    centro_depto = ct["B5"].value  # centro "do departamento" na base (ADM01)
    eventos, usados = [], set()
    for r in range(2, csv_v.max_row + 1):
        descricao = csv_v.cell(r, 2).value
        if descricao is None:
            continue
        ref = re.match(r"=Contabil!F(\d+)$", str(csv_f.cell(r, 2).value))
        linha = int(ref.group(1)) if ref else None
        calculo = CALCULOS.get(descricao)
        if linha:
            usados.add(linha)
        centro = csv_v.cell(r, 4).value
        cod_e = ct.cell(linha, 5).value if linha else None
        eventos.append({
            "ordem": len(eventos) + 1,
            "codigo_pdf": f"{int(cod_e):03d}" if cod_e not in (None, "") else None,
            "nome_pdf": None if calculo else limpar_nome(descricao),
            "descricao": descricao,
            "secao": "calculado" if calculo else secao_da_linha(linha),
            "tipo_lancamento": csv_v.cell(r, 3).value,
            "centro_custo_fixo": None if centro == centro_depto else centro,
            "item_despesa": inteiro(csv_v.cell(r, 6).value),
            "gera_lancamento": True,
            "calculo": calculo,
        })

    # ---- eventos conhecidos que NÃO geram lançamento (não disparam alerta de "evento novo") ----
    for r in list(ADICIONAIS) + list(DESCONTOS) + list(TOTAIS) + list(GPS) + list(PIS) + list(FGTS):
        nome = ct.cell(r, 6).value
        if r in usados or not nome or (r in GPS and r == 52):
            continue
        cod_e = ct.cell(r, 5).value
        eventos.append({
            "ordem": 1000 + r,
            "codigo_pdf": f"{int(cod_e):03d}" if isinstance(cod_e, (int, float)) else None,
            "nome_pdf": limpar_nome(nome),
            "descricao": limpar_nome(nome),
            "secao": secao_da_linha(r),
            "tipo_lancamento": None,
            "centro_custo_fixo": None,
            "item_despesa": None,
            "gera_lancamento": False,
            "calculo": None,
        })

    SAIDA.mkdir(parents=True, exist_ok=True)
    for nome, dados in (("centros_custo", centros), ("itens_despesa", itens), ("eventos", eventos)):
        (SAIDA / f"{nome}.json").write_text(json.dumps(dados, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    lanca = sum(e["gera_lancamento"] for e in eventos)
    print(f"{len(centros)} centros de custo, {len(itens)} itens de despesa, {len(eventos)} eventos ({lanca} geram lançamento)")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else PADRAO)
