import math
import openpyxl
from quinzena.extrato import carregar_extrato

HEAD = ["Data", "Banco", "Origem", "Tipo", "Título", "Descrição", "Centro de Custo",
        "Item de Despesa", "Fonte de Recurso", "Valor", "Saldo", "Obs"]


def _xlsx(tmp_path, linhas):
    wb = openpyxl.Workbook(); ws = wb.active
    ws["B1"] = "DDATA - SOL"
    for j, h in enumerate(HEAD):
        ws.cell(2, 2 + j, h)
    ws["L3"] = 70312.29
    r = 4
    for lin in linhas:
        for j, v in enumerate(lin):
            ws.cell(r, 2 + j, v)
        r += 1
    ws.cell(r + 1, 11, "EXTRATOBANCO--->>>")
    ws.cell(r + 1, 12, 177956.86)
    p = tmp_path / "e.xlsx"; wb.save(p); return str(p)


def _lin(**k):
    b = dict(d="01/04/2026", b="M", o="Lançamento", t="Debito", ti="NFe - 1", de="Pag", c="CTPAS",
             i="520 - X", f="—", v=-10, s=0, ob="x")
    b.update(k)
    return [b["d"], b["b"], b["o"], b["t"], b["ti"], b["de"], b["c"], b["i"], b["f"], b["v"], b["s"], b["ob"]]


def test_ignora_rodape_e_linhas_vazias(tmp_path):
    df = carregar_extrato(_xlsx(tmp_path, [_lin(), _lin(ti="NFe - 2")]))
    assert len(df) == 2
    assert list(df.titulo) == ["NFe - 1", "NFe - 2"]


def test_traco_vira_none(tmp_path):
    df = carregar_extrato(_xlsx(tmp_path, [_lin(i="—")]))
    assert df.item.iloc[0] is None and df.fonte.iloc[0] is None


def test_valor_invalido_vira_nan(tmp_path):
    df = carregar_extrato(_xlsx(tmp_path, [_lin(v="abc")]))
    assert math.isnan(df.valor.iloc[0])
