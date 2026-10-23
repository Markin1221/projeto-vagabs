import json
import openpyxl
import pandas as pd
from quinzena.revisao import exportar_revisao, importar_decisoes


def _df(mk):
    df = mk([dict(titulo="A", valor=-10.0), dict(titulo="B NOVO 1", valor=-20.0), dict(titulo="C", valor=100.0, tipo="Credito")])
    df["sug_centro"] = ["C", None, None]; df["sug_item"] = ["520 - X", None, None]
    df["sug_fonte"] = [None, None, None]; df["origem_regra"] = ["sol", "manual", "manual"]
    df["manual"] = [False, True, True]
    return df


def test_manual_aparece_primeiro(mk, tmp_path):
    p = str(tmp_path / "r.xlsx"); exportar_revisao(_df(mk), p)
    rev = pd.read_excel(p, sheet_name="Revisão")
    assert list(rev.manual[:2]) == [True, True] and rev.manual.iloc[2] == False


def test_soma_resumo_igual_soma_extrato(mk, tmp_path):
    df = _df(mk); p = str(tmp_path / "r.xlsx"); exportar_revisao(df, p)
    ws = openpyxl.load_workbook(p)["Resumo"]
    achados = {r[0].value: r[2].value for r in ws.iter_rows() if r[0].value in ("TOTAL", "TOTAL EXTRATO", "DIFERENÇA")}
    assert round(achados["TOTAL"], 2) == round(df.valor.sum(), 2) == round(achados["TOTAL EXTRATO"], 2)
    assert achados["DIFERENÇA"] == 0


def test_importar_decisoes_cria_regra_e_retorna_contagem(mk, tmp_path):
    p = str(tmp_path / "r.xlsx"); exportar_revisao(_df(mk), p)
    wb = openpyxl.load_workbook(p); ws = wb["Revisão"]
    cols = {c.value: c.column for c in ws[1]}
    ws.cell(2, cols["decisao_item"], "580 - EMPRESTIMO AO PERMISSIONÁRIO")
    wb.save(p)
    rp = tmp_path / "regras.json"; rp.write_text("[]")
    assert importar_decisoes(p, str(rp)) == 1
    regras = json.loads(rp.read_text(encoding="utf-8"))
    assert regras[0]["item"].startswith("580") and regras[0]["centro"].startswith("CTPAS")


def test_importar_sem_decisoes_retorna_zero(mk, tmp_path):
    p = str(tmp_path / "r.xlsx"); exportar_revisao(_df(mk), p)
    rp = tmp_path / "regras.json"; rp.write_text("[]")
    assert importar_decisoes(p, str(rp)) == 0
