"""Planilha de revisão (MANUAL no topo) e importação das decisões como regras."""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from .classificador import CENTRO_PADRAO
from .padroes import chave_favorecido
from .regras import Regra, salvar_regra

COLS_REV = ["linha_origem", "data", "tipo", "titulo", "descricao", "valor", "sug_centro", "sug_item",
            "sug_fonte", "origem_regra", "manual", "decisao_centro", "decisao_item", "decisao_fonte"]


def exportar_revisao(df: pd.DataFrame, caminho: str) -> None:
    rev = df.copy()
    for c in ("decisao_centro", "decisao_item", "decisao_fonte"):
        rev[c] = None
    rev = rev.sort_values("manual", ascending=False, kind="stable")[COLS_REV]
    g = df.assign(grupo=df["sug_item"].fillna("MANUAL")).groupby("grupo")["valor"].agg(["count", "sum"]).reset_index()
    total = float(df["valor"].sum())
    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        rev.to_excel(w, sheet_name="Revisão", index=False)
        ws = w.book.create_sheet("Resumo")
        ws.append(["item", "qtd", "soma"])
        for _, r in g.iterrows():
            ws.append([r["grupo"], int(r["count"]), round(float(r["sum"]), 2)])
        soma = round(float(g["sum"].sum()), 2)
        ws.append(["TOTAL", int(g["count"].sum()), soma])
        ws.append(["TOTAL EXTRATO", len(df), round(total, 2)])
        ws.append(["DIFERENÇA", None, round(soma - total, 2)])
        amarelo = PatternFill("solid", fgColor="FFFF00")
        wr = w.book["Revisão"]
        for i, m in enumerate(rev["manual"], start=2):
            if m:
                for c in wr[i]:
                    c.fill = amarelo


def importar_decisoes(caminho: str, regras_path: str) -> int:
    rev = pd.read_excel(caminho, sheet_name="Revisão")
    n = 0
    for _, r in rev.iterrows():
        item = r.get("decisao_item")
        if pd.isna(item) or not str(item).strip():
            continue
        centro = r.get("decisao_centro")
        centro = CENTRO_PADRAO if pd.isna(centro) or not str(centro).strip() else str(centro)
        kw = chave_favorecido(None if pd.isna(r["descricao"]) else r["descricao"]) \
            if str(r.get("tipo")) == "Debito" else None
        kw = kw or (None if pd.isna(r["titulo"]) else str(r["titulo"]).upper())
        if not kw:
            continue
        salvar_regra(regras_path, Regra(kw, centro, str(item).strip(), "", "" if pd.isna(r["titulo"]) else str(r["titulo"])))
        n += 1
    return n
