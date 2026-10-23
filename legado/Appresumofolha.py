import calendar
import re
import sys
import unicodedata
from copy import copy
from datetime import datetime

import openpyxl
import pandas as pd
import pdfplumber
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from PyQt6.QtWidgets import (
    QApplication, QWidget, QLabel, QPushButton, QVBoxLayout, QFileDialog, QTextEdit, QInputDialog
)

NUM = re.compile(r"^\d{1,3}(\.\d{3})*,\d{2}$|^\d+$")
NUM_BR = re.compile(r"^\d{1,3}(\.\d{3})*,\d{2}$")
PERIODO = r"\d{2}/\d{2}/\d{2} a \d{2}/\d{2}/\d{2}"


def br_para_float(txt):
    """'54.244,65' -> 54244.65 ; '' -> None"""
    if txt is None or str(txt).strip() == "":
        return None
    return float(str(txt).replace(".", "").replace(",", "."))


def agrupar_linhas(words, tol=3):
    """Agrupa palavras na mesma linha (por 'top') e ordena da esquerda p/ direita."""
    linhas = []
    for w in sorted(words, key=lambda w: w["top"]):
        if linhas and abs(linhas[-1][0] - w["top"]) <= tol:
            linhas[-1][1].append(w)
        else:
            linhas.append([w["top"], [w]])
    return [sorted(l[1], key=lambda w: w["x0"]) for l in linhas]


# ---------- Localização das partes (por conteúdo, não por página) ----------
def y_blocos(words):
    """Topo da faixa GPS/FGTS/DARF (título 'DARF PIS'); None se a página não tem."""
    for i, w in enumerate(words[:-1]):
        if w["text"] == "DARF" and words[i + 1]["text"] == "PIS":
            return w["top"] - 2
    return None


# ---------- Tabela de adicionais / descontos ----------
def ler_tabela(page):
    """Colunas de valores vêm dos títulos (Ativos/Demitidos/... ou Sócios/Autônomos)."""
    words = page.extract_words()
    y_fim = y_blocos(words) or page.height
    linhas_pdf = agrupar_linhas([w for w in words if w["top"] < y_fim])

    pos = next((i for i, l in enumerate(linhas_pdf) if l[0]["text"] == "ADICIONAIS"), None)
    if pos is None:
        return None
    cab = linhas_pdf[pos]
    idx = next((i for i, w in enumerate(cab) if w["text"] == "DESCONTOS"), None)
    if idx is None:
        return None

    resto, x_col, k = cab[idx + 1:], {}, 0
    while k < len(resto):
        if resto[k]["text"] == "Valores" and k + 1 < len(resto):  # "Valores Sócios"
            x_col[resto[k + 1]["text"].capitalize()] = resto[k + 1]["x1"]
            k += 2
        else:
            x_col[resto[k]["text"].capitalize()] = resto[k]["x1"]
            k += 1
    if not x_col:
        return None

    linhas, tipo = [], "Adicional"
    for lin in linhas_pdf[pos + 1:]:
        textos = [w["text"] for w in lin]
        if not any(NUM.match(t) for t in textos):
            continue
        codigo = textos[0] if re.fullmatch(r"\d{3}", textos[0]) else ""
        valores = {c: None for c in x_col}
        nome = []
        for w in lin[1 if codigo else 0:]:
            col = min(x_col, key=lambda c: abs(x_col[c] - w["x1"]))
            if NUM.match(w["text"]) and abs(x_col[col] - w["x1"]) < 25:
                valores[col] = br_para_float(w["text"])
            else:
                nome.append(w["text"])
        nome = " ".join(nome)
        eh_total_grupo = nome.startswith("TOTAL DE ADICIONAIS") or nome.startswith("TOTAL DE DESCONTOS")
        linhas.append({"Tipo": tipo if (codigo or eh_total_grupo) else "",
                       "Código": codigo, "Nome": nome, **valores})
        if nome.startswith("TOTAL DE ADICIONAIS"):
            tipo = "Desconto"
    return pd.DataFrame(linhas)


# ---------- Blocos GPS / FGTS / DARF PIS ----------
def ler_bloco(page, x0, x1, y0, y1, secao_inicial=""):
    """Lê um bloco 'Nome : valor'. Linha sem valor vira cabeçalho de seção."""
    area = page.crop((x0, y0, x1, y1))
    dados, secao = [], secao_inicial
    for lin in agrupar_linhas(area.extract_words()):
        textos = [w["text"] for w in lin]
        if len(textos) > 1 and NUM_BR.match(textos[-1]):
            nome = " ".join(textos[:-1]).strip().rstrip(":").strip()
            nome = re.sub(r"=(\S)", r"= \1", nome)
            dados.append({"Seção": secao, "Nome": nome, "Valor": br_para_float(textos[-1])})
        else:
            secao = " ".join(textos)
    return pd.DataFrame(dados)


def ler_blocos(page):
    words = page.extract_words()
    y0 = y_blocos(words)
    if y0 is None:
        return None
    larg = page.width

    # fim da faixa: antes de "INFORMAÇÕES AUXILIARES"
    y1 = next((w["top"] - 2 for w in words if w["text"].upper().startswith("AUXILIARES")), page.height)

    # início do FGTS = primeira palavra "BASES"; início do DARF PIS = "Base" seguido de "PIS"
    x_fgts = next((w["x0"] for w in words if w["text"] == "BASES" and w["top"] > y0), larg / 3) - 3
    x_darf = larg * 2 / 3
    for i, w in enumerate(words[:-1]):
        if w["text"] == "Base" and words[i + 1]["text"] == "PIS" and w["top"] > y0:
            x_darf = w["x0"] - 3
            break

    gps = ler_bloco(page, 0, x_fgts, y0, y1)
    fgts = ler_bloco(page, x_fgts, x_darf, y0, y1)
    darf = ler_bloco(page, x_darf, larg, y0, y1)
    return {
        "GPS": gps[["Nome", "Valor"]],
        "FGTS": fgts[fgts["Seção"] != "F G T S"].reset_index(drop=True),
        "DARF PIS": darf.reset_index(drop=True),
    }


# ---------- Férias (só existe em alguns relatórios) ----------
def ler_ferias(page):
    padrao = re.compile(
        rf"^(\d{{6}})\s+(.+?)\s+(\d{{1,3}}(?:\.\d{{3}})*,\d{{2}})\s+(\d{{2}}/\d{{2}}/\d{{2}})\s+({PERIODO})\s+({PERIODO})"
    )
    dados = []
    for linha in (page.extract_text() or "").splitlines():
        m = padrao.match(linha.strip())
        if m:
            dados.append({
                "Código": m.group(1), "Funcionário": m.group(2),
                "Valor líquido": br_para_float(m.group(3)), "Data pagamento": m.group(4),
                "Período de férias": m.group(5), "Período aquisitivo": m.group(6),
            })
    return pd.DataFrame(dados)


# ---------- Validação ----------
def validar(df1):
    avisos = []
    cols = [c for c in df1.columns if c not in ("Tipo", "Código", "Nome")]
    itens = df1[df1["Código"] != ""]

    def total(nome):
        r = df1[df1["Nome"] == nome]
        return r.iloc[0] if not r.empty else None

    t_adic, t_desc, liq = total("TOTAL DE ADICIONAIS"), total("TOTAL DE DESCONTOS"), total("TOTAL LÍQUIDO A PAGAR")
    for col in cols:
        s_adic = itens[itens["Tipo"] == "Adicional"][col].fillna(0).sum()
        s_desc = itens[itens["Tipo"] == "Desconto"][col].fillna(0).sum()
        if t_adic is not None and pd.notna(t_adic[col]) and abs(s_adic - t_adic[col]) > 0.01:
            avisos.append(f"{col}: soma dos adicionais ({s_adic:.2f}) difere do total ({t_adic[col]})")
        if t_desc is not None and pd.notna(t_desc[col]) and abs(s_desc - t_desc[col]) > 0.01:
            avisos.append(f"{col}: soma dos descontos ({s_desc:.2f}) difere do total ({t_desc[col]})")
        if liq is not None and pd.notna(liq[col]) and abs(s_adic - s_desc - liq[col]) > 0.01:
            avisos.append(f"{col}: adicionais - descontos ({s_adic - s_desc:.2f}) difere do líquido ({liq[col]})")
    return avisos


# ---------- Avaliador de fórmulas (para gravar a aba de lançamentos só com valores) ----------
class _X:
    """Valor de célula com as regras do Excel: vazio vale 0 na conta e '' no texto."""
    __slots__ = ("v",)

    def __init__(self, v):
        self.v = v

    def _n(self):
        v = self.v
        if v is None or v == "":
            return 0.0
        if isinstance(v, (bool, int, float)):
            return float(v)
        try:  # texto que parece número ('75') vale como número na conta, igual ao Excel
            return float(v)
        except ValueError:
            raise TypeError(f"valor não numérico: {v!r}")

    def _s(self):
        v = self.v
        if v is None:
            return ""
        if isinstance(v, float) and v == int(v):
            return str(int(v))
        return str(v)

    def __add__(s, o): return _X(s._n() + _x(o)._n())
    def __sub__(s, o): return _X(s._n() - _x(o)._n())
    def __mul__(s, o): return _X(s._n() * _x(o)._n())
    def __pow__(s, o): return _X(s._n() ** _x(o)._n())
    def __neg__(s): return _X(-s._n())
    def __pos__(s): return s
    def __or__(s, o): return _X(s._s() + _x(o)._s())  # '&' do Excel

    def __truediv__(s, o):
        d = _x(o)._n()
        return _X("#DIV/0!") if d == 0 else _X(s._n() / d)

    def __eq__(s, o):
        a, b = s.v, _x(o).v
        if isinstance(a, str) or isinstance(b, str):
            return _X(str("" if a is None else a).casefold() == str("" if b is None else b).casefold())
        return _X(_X(a)._n() == _X(b)._n())

    def __ne__(s, o): return _X(not (s == o).v)
    def __lt__(s, o): return _X(s._n() < _x(o)._n())
    def __gt__(s, o): return _X(s._n() > _x(o)._n())
    def __le__(s, o): return _X(s._n() <= _x(o)._n())
    def __ge__(s, o): return _X(s._n() >= _x(o)._n())


def _x(v):
    return v if isinstance(v, _X) else _X(v)


class AvaliadorFormulas:
    """Calcula o valor das células de um workbook openpyxl (funções: SUM, SUMIFS, IF)."""

    def __init__(self, wb):
        self.wb = wb
        self.cache = {}
        self.pilha = set()

    def valor(self, folha, coord):
        chave = (folha, coord)
        if chave in self.cache:
            return self.cache[chave]
        if chave in self.pilha:
            raise ValueError(f"referência circular em {folha}!{coord}")
        v = self.wb[folha][coord].value
        if isinstance(v, str) and v.startswith("="):
            self.pilha.add(chave)
            try:
                v = self._calcular(folha, v)
            finally:
                self.pilha.discard(chave)
        self.cache[chave] = v
        return v

    def _calcular(self, folha, formula):
        codigo = self._traduzir(folha, formula)
        ns = {"X": _X, "C": self._c, "R": self._r, "SUM": self._sum, "SUMIFS": self._sumifs, "IF": self._if}
        return eval(codigo, {"__builtins__": {}}, ns).v

    def _traduzir(self, folha, formula):
        partes = []
        for t in Tokenizer(formula).items:
            if t.type == Token.WSPACE:
                continue
            if t.type == Token.OPERAND:
                if t.subtype == Token.RANGE:
                    ref = t.value
                    sh, addr = ref.rsplit("!", 1) if "!" in ref else (folha, ref)
                    sh = sh.strip("'").replace("''", "'")
                    addr = addr.replace("$", "")
                    partes.append(f'R({sh!r},{addr!r})' if ":" in addr else f'C({sh!r},{addr!r})')
                elif t.subtype == Token.NUMBER:
                    partes.append(f"X({float(t.value)!r})")
                elif t.subtype == Token.TEXT:
                    partes.append(f"X({t.value[1:-1].replace(chr(34) * 2, chr(34))!r})")
                else:  # lógico ou erro (#REF!)
                    partes.append(f"X({t.value!r})")
            elif t.type == Token.FUNC:
                partes.append(t.value.upper() if t.subtype == Token.OPEN else ")")
            elif t.type in (Token.SEP, Token.PAREN):
                partes.append(t.value)
            elif t.type in (Token.OP_IN, Token.OP_PRE):
                partes.append({"&": "|", "=": "==", "<>": "!=", "^": "**"}.get(t.value, t.value))
            elif t.type == Token.OP_POST:  # porcentagem
                partes.append("*X(0.01)")
        return "".join(partes)

    def _c(self, folha, coord):
        return _X(self.valor(folha, coord))

    def _r(self, folha, addr):
        c1, r1, c2, r2 = range_boundaries(addr)
        return [self.valor(folha, f"{get_column_letter(c)}{r}")
                for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    @staticmethod
    def _sum(*args):
        total = 0.0
        for a in args:
            for v in (a if isinstance(a, list) else [a.v]):
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    total += v
        return _X(total)

    @staticmethod
    def _sumifs(soma, criterio_rng, criterio):
        crit, total = criterio.v, 0.0
        for s, c in zip(soma, criterio_rng):
            if isinstance(crit, str):
                ok = isinstance(c, str) and c.casefold() == crit.casefold()
            else:
                ok = isinstance(c, (int, float)) and c == crit
            if ok and isinstance(s, (int, float)) and not isinstance(s, bool):
                total += s
        return _X(total)

    @staticmethod
    def _if(cond, a, b=None):
        v = cond.v
        return a if (v and v != "#DIV/0!") else (b if b is not None else _X(False))


def congelar_formulas(wb, folha):
    """Troca as fórmulas da aba por seus valores calculados (valores de dinheiro com 2 casas)."""
    av = AvaliadorFormulas(wb)
    ws = wb[folha]
    novos = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                v = av.valor(folha, cell.coordinate)
                if isinstance(v, float):
                    v = round(v, 2) if cell.column_letter == "L" else (int(v) if v == int(v) else v)
                if isinstance(v, str) and v.startswith("#"):
                    v = None
                novos[cell.coordinate] = v
    for coord, v in novos.items():
        ws[coord].value = v
    return len(novos)


# ---------- Preenchimento da planilha base ----------
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho",
         "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


def _norm(s):
    s = unicodedata.normalize("NFD", str(s).casefold())
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip()


def ler_cabecalho(pdf_path):
    """(data_inicio, data_fim, nome_do_departamento) lidos do cabeçalho do relatório."""
    with pdfplumber.open(pdf_path) as pdf:
        texto = pdf.pages[0].extract_text() or ""
    m = re.search(r"Per.odo\s*:\s*(\d{2}/\d{2}/\d{4})\s+a\s+(\d{2}/\d{2}/\d{4})", texto)
    d = re.search(r"Departamento\s*:\s*(?:\d+\s+)?(.+)", texto)
    if not m or not d:
        raise ValueError("Não achei Período/Departamento no cabeçalho do PDF")
    ini, fim = (datetime.strptime(x, "%d/%m/%Y") for x in m.groups())
    return ini, fim, d.group(1).strip()


def limpar_e_escrever(ws, colunas, linhas):
    """Apaga os valores das colunas (a partir da linha 2, mantendo formato) e escreve as linhas novas."""
    estilos = {c: copy(ws.cell(2, c)._style) for c in colunas}
    for r in range(2, ws.max_row + 1):
        for c in colunas:
            ws.cell(r, c).value = None
    for i, linha in enumerate(linhas, start=2):
        for c, v in zip(colunas, linha):
            cell = ws.cell(i, c)
            cell.value = None if (v is None or (isinstance(v, float) and pd.isna(v)) or v == "") else v
            cell._style = copy(estilos[c])


def preencher_base(pdf_path, base_path, saida, escolher_centro=None):
    """Abre a planilha base, troca os valores pelos do PDF e salva em 'saida'. Devolve lista de avisos."""
    avisos = []
    ini, fim, depto_pdf = ler_cabecalho(pdf_path)

    df1 = blocos = None
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            if df1 is None:
                df1 = ler_tabela(page)
            if blocos is None:
                blocos = ler_blocos(page)
    if df1 is None or blocos is None:
        raise ValueError("Não encontrei a tabela de adicionais/descontos ou os blocos GPS/FGTS/DARF PIS")
    avisos += validar(df1)

    # nomes iguais aos da base: tira o "Parc. 6/48" do empréstimo
    df1["Nome"] = df1["Nome"].str.replace(r"\s+Parc\.\s*\d+/\d+\s*$", "", regex=True)

    wb = openpyxl.load_workbook(base_path)

    # ---- centro de custo: procura o departamento na tabela da própria base ----
    centros = []  # (codigo, nome)
    for r in range(2, 60):
        cod, nome = (wb["Itens Despesas Resumido"].cell(r, c).value for c in (7, 8))
        if cod and nome:
            centros.append((str(cod), str(nome)))
    alvo = _norm(depto_pdf)
    achado = next((c for c in centros if _norm(c[1]) == alvo), None) \
        or next((c for c in centros if alvo in _norm(c[1]) or _norm(c[1]) in alvo), None)
    if achado:
        centro, nome_depto = achado
    else:
        nome_depto = depto_pdf.title()
        centro = escolher_centro(depto_pdf, centros) if escolher_centro else None
        if not centro:
            centro = wb["Contabil"]["B5"].value
            avisos.append(f"Departamento '{depto_pdf}' não está na lista de centros de custo; mantive {centro}.")

    # ---- Resumo Folha ----
    cols = [c for c in df1.columns if c not in ("Tipo", "Código", "Nome")]
    ws = wb["Resumo Folha"]
    if "Total" in cols:
        outras = [c for c in cols if c != "Total"]
        total = df1["Total"]
    else:
        outras = cols
        total = df1[cols].sum(axis=1, min_count=1)
    outras = outras[:3]
    linhas = []
    for i, r in df1.iterrows():
        vals = [r[c] for c in outras] + [None] * (3 - len(outras)) + [total[i]]
        linhas.append([r["Tipo"], r["Código"], r["Nome"]] + vals)
    limpar_e_escrever(ws, range(1, 8), [])  # só limpa; cabeçalho é reescrito abaixo
    for j, nome in enumerate(outras + [None] * (3 - len(outras)) + ["Total"], start=4):
        ws.cell(1, j).value = nome
    limpar_e_escrever(ws, range(1, 8), linhas)

    # ---- GPS / FGTS / DARF PIS (só as colunas de dados; fórmulas ao lado ficam) ----
    limpar_e_escrever(wb["GPS"], [1, 2], blocos["GPS"][["Nome", "Valor"]].values.tolist())
    for nome in ("FGTS", "DARF PIS"):
        limpar_e_escrever(wb[nome], [1, 2, 3], blocos[nome][["Seção", "Nome", "Valor"]].values.tolist())

    # ---- lancamentos CSV: título, data e observação ----
    dia = min(30, calendar.monthrange(fim.year, fim.month)[1])
    data_pag = datetime(fim.year, fim.month, dia)
    lc = wb["lancamentos CSV"]
    lc["A2"] = f"Folha {nome_depto}"
    lc["J2"] = data_pag
    lc["M2"] = f"Folha {MESES[fim.month - 1]} {fim.year % 100:02d} {nome_depto}"

    # ---- Contabil: título e centro de custo ----
    ct = wb["Contabil"]
    ct["B1"] = f"Planilha Lançamentos Contábil - {nome_depto}"
    ct["B5"] = centro

    # ---- confere se todo evento do PDF tem linha na aba Contabil ----
    nomes_contabil = {_norm(ct.cell(r, 6).value) for r in range(1, 130) if ct.cell(r, 6).value}
    for _, r in df1[df1["Código"] != ""].iterrows():
        valor = r[outras[0]] if outras else None
        if _norm(r["Nome"]) not in nomes_contabil and valor not in (None, 0) and not pd.isna(valor):
            avisos.append(f"Evento do PDF sem linha na aba Contabil (valor ignorado): {r['Código']} {r['Nome']} = {valor:.2f}")

    # aba de lançamentos: só valores (sem fórmulas); as demais abas continuam com fórmulas
    congelar_formulas(wb, "lancamentos CSV")
    wb.calculation.fullCalcOnLoad = True  # o Excel recalcula as fórmulas das outras abas ao abrir
    wb.save(saida)
    avisos.insert(0, f"{nome_depto} ({centro}) - {MESES[fim.month - 1]}/{fim.year} - data do lançamento {data_pag:%d/%m/%Y}")
    return avisos


# ---------- Interface ----------
class AppResumoFolha(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Resumo Geral da Folha (PDF) -> Planilha base")
        self.setGeometry(400, 200, 650, 450)
        self.arquivo_pdf = ""
        self.arquivo_base = ""

        self.layout = QVBoxLayout()
        self.btn_base = QPushButton("1) Selecionar PLANILHA BASE (.xlsx)")
        self.btn_base.clicked.connect(self.selecionar_base)
        self.layout.addWidget(self.btn_base)
        self.label_base = QLabel("Nenhum arquivo selecionado")
        self.layout.addWidget(self.label_base)

        self.btn_pdf = QPushButton("2) Selecionar PDF do Resumo Geral")
        self.btn_pdf.clicked.connect(self.selecionar_pdf)
        self.layout.addWidget(self.btn_pdf)
        self.label_pdf = QLabel("Nenhum arquivo selecionado")
        self.layout.addWidget(self.label_pdf)

        self.btn_gerar = QPushButton("3) Preencher base e salvar")
        self.btn_gerar.clicked.connect(self.processar)
        self.layout.addWidget(self.btn_gerar)

        self.texto_status = QTextEdit()
        self.texto_status.setReadOnly(True)
        self.layout.addWidget(self.texto_status)
        self.setLayout(self.layout)

    def selecionar_base(self):
        arquivo, _ = QFileDialog.getOpenFileName(self, "Selecione a planilha base", "", "Arquivos Excel (*.xlsx)")
        if arquivo:
            self.arquivo_base = arquivo
            self.label_base.setText(arquivo)

    def selecionar_pdf(self):
        arquivo, _ = QFileDialog.getOpenFileName(self, "Selecione o PDF", "", "Arquivos PDF (*.pdf)")
        if arquivo:
            self.arquivo_pdf = arquivo
            self.label_pdf.setText(arquivo)

    def perguntar_centro(self, depto, centros):
        opcoes = [f"{cod} - {nome}" for cod, nome in centros]
        escolha, ok = QInputDialog.getItem(
            self, "Centro de custo", f"Não reconheci o departamento '{depto}'.\nEscolha o centro de custo:",
            opcoes, 0, False)
        return escolha.split(" - ")[0] if ok else None

    def processar(self):
        if not self.arquivo_base or not self.arquivo_pdf:
            self.texto_status.append("Selecione a planilha base e o PDF antes de gerar.")
            return
        try:
            ini, fim, depto = ler_cabecalho(self.arquivo_pdf)
        except Exception as erro:
            self.texto_status.append(f"Erro ao ler o PDF: {erro}")
            return
        sugestao = f"Contabil_folha_{depto.title().replace(' ', '_')}_{fim:%m-%Y}.xlsx"
        caminho, _ = QFileDialog.getSaveFileName(self, "Salvar planilha preenchida", sugestao, "Arquivos Excel (*.xlsx)")
        if not caminho:
            self.texto_status.append("Salvamento cancelado.")
            return
        try:
            avisos = preencher_base(self.arquivo_pdf, self.arquivo_base, caminho, self.perguntar_centro)
        except Exception as erro:
            self.texto_status.append(f"Erro: {erro}")
            return
        self.texto_status.append(avisos[0])
        if len(avisos) > 1:
            self.texto_status.append("Atenção:")
            for a in avisos[1:]:
                self.texto_status.append(f" - {a}")
        else:
            self.texto_status.append("Validação OK (totais conferem).")
        self.texto_status.append(f"Arquivo salvo em: {caminho}")


if __name__ == "__main__":
    app = QApplication(sys.argv)
    import estilo; estilo.aplicar(app)
    janela = AppResumoFolha()
    janela.show()
    sys.exit(app.exec())